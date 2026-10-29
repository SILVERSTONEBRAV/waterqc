from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import zipfile
from datetime import date, timedelta

from django.core.exceptions import PermissionDenied
from django.core.cache import cache
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from django.views import View

from core.mixins import WanjalaMixin
from qc.services.advanced_reports import (
    PERIOD_LENGTHS,
    SORT_FIELDS,
    WATER_TYPES,
    build_analytics,
    count_analysis_records,
    daily_parameter_summaries,
    image_report,
    json_safe,
    load_result_rows,
    multivariate_analysis,
    parameter_options,
    parameter_summaries,
    phase_options,
    rows_for_export,
    sample_summary,
)
from t.models import TreatmentPlant


logger = logging.getLogger(__name__)
REPORT_SCHEMA_VERSION = "2.0"


def cached_result_rows(**scope):
    """Short-lived cache for the normalized, plant-scoped result fact rows."""
    serialized = json.dumps(
        {key: str(value) if value is not None else "" for key, value in sorted(scope.items())},
        sort_keys=True,
    )
    digest = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
    key = f"qcr:result-rows:v2:{digest}"
    cached = cache.get(key)
    if cached is not None:
        return cached
    rows = load_result_rows(**scope)
    cache.set(key, rows, timeout=600)
    return rows


PERIOD_CHOICES = (
    ("day", "Today"),
    ("week", "Last 7 days"),
    ("month", "Last 30 days"),
    ("quarter", "Last 13 weeks"),
    ("half_year", "Last 6 months"),
    ("year", "Last 12 months"),
    ("custom", "Custom range"),
)

GROUP_CHOICES = (
    ("day", "Daily"),
    ("week", "Weekly"),
    ("month", "Monthly"),
    ("quarter", "Quarterly"),
    ("half_year", "Semiannual"),
    ("year", "Annual"),
)

REPORT_TYPE_PERMISSIONS = {
    "raw": "qc.view_wateranalysis",
    "coagulated": "qc.view_wateranalysis",
    "clarified": "qc.view_wateranalysis",
    "treated": "qc.view_wateranalysis",
    "filtered_water": "filtered_water.view_filteredwateranalysis",
    "microbial": "microbial.view_microbialanalysis",
    "heavy_metal": "heavy_metal.view_heavymetalanalysis",
}


class QCReportsBase(WanjalaMixin, View):
    """Plant-scoped, parameter-aware QC reporting and analytics workspace."""

    template_name = "qc/reports/dashboard.html"
    permission_required_view = None
    raise_exception = True

    def dispatch(self, request, *args, **kwargs):
        self.request = request
        self.action = "view"
        # WanjalaMixin performs login and lifecycle checks in dispatch. Resolve
        # authentication and the report permission before it can handle exports.
        if not request.user.is_authenticated:
            return self.handle_no_permission()
        if not self.permitted_water_types(request.user):
            raise PermissionDenied("You do not have permission to view QC reports.")
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, *args, **kwargs):
        data = self.get_report_data(request)
        export_type = (request.GET.get("export") or "").strip().lower()
        if export_type:
            return self.export_report(export_type, data)
        return render(request, self.template_name, self.template_context(data))

    def handle_export(self, request):
        """Use WanjalaMixin's export dispatch while supporting analytics formats."""
        export_type = (request.GET.get("export") or request.GET.get("type") or "").strip().lower()
        return self.export_report(export_type, self.get_report_data(request))

    @staticmethod
    def download(content, filename, content_type):
        response = HttpResponse(content, content_type=content_type)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        return response

    def selected_plant(self, request):
        user = request.user
        if user.is_superuser:
            available_plants = TreatmentPlant.objects.order_by("name")
            plant_id = request.GET.get("plant") or request.session.get("qc_selected_plant")
            try:
                plant = available_plants.filter(pk=plant_id).first() if plant_id else available_plants.first()
            except (TypeError, ValueError):
                plant = available_plants.first()
            if plant:
                request.session["qc_selected_plant"] = plant.pk
        else:
            plant = getattr(user, "assigned_treatment_plant", None)
            available_plants = [plant] if plant else []
        return plant, list(available_plants)

    @staticmethod
    def permitted_water_types(user):
        if user.is_superuser:
            return WATER_TYPES
        return tuple(
            (key, label) for key, label in WATER_TYPES
            if user.has_perm(REPORT_TYPE_PERMISSIONS.get(key, ""))
        )

    @staticmethod
    def parse_date(value, default):
        try:
            return date.fromisoformat(value) if value else default
        except (TypeError, ValueError):
            return default

    def get_report_data(self, request, include_full_analytics=False):
        today = timezone.localdate()
        period = request.GET.get("period", "month")
        if period not in dict(PERIOD_CHOICES):
            period = "month"
        end_date = self.parse_date(request.GET.get("end"), today)
        if period in PERIOD_LENGTHS:
            start_date = end_date - timedelta(days=PERIOD_LENGTHS[period] - 1)
        else:
            start_date = self.parse_date(request.GET.get("start"), end_date - timedelta(days=29))
        if start_date > end_date:
            start_date, end_date = end_date, start_date

        available_water_types = self.permitted_water_types(request.user)
        allowed_water_type_keys = {value for value, _ in available_water_types}
        water_type = request.GET.get("water_type", "raw")
        if water_type not in allowed_water_type_keys:
            water_type = available_water_types[0][0] if available_water_types else "raw"
        group_by = request.GET.get("group_by", "day")
        if group_by not in dict(GROUP_CHOICES):
            group_by = "day"

        plant, available_plants = self.selected_plant(request)
        parameters = parameter_options(plant, water_type)
        parameter_names = {item["name"] for item in parameters}
        parameter_name = request.GET.get("parameter", "").strip()
        if parameter_name not in parameter_names:
            parameter_name = ""
        focus_parameter = request.GET.get("focus", "").strip()
        if focus_parameter not in parameter_names:
            focus_parameter = ""
        if parameter_name:
            focus_parameter = parameter_name
        phases = phase_options(plant) if water_type not in {"raw"} else []
        valid_phase_ids = {str(item["id"]) for item in phases}
        phase_id = request.GET.get("phase", "")
        if phase_id not in valid_phase_ids:
            phase_id = ""

        search = request.GET.get("search", "").strip()[:120]
        if plant:
            row_scope = dict(
                plant=plant,
                start_date=start_date,
                end_date=end_date,
                water_type=water_type,
                parameter_name=parameter_name,
                phase_id=phase_id,
                search=search,
            )
            rows = cached_result_rows(**row_scope)
        else:
            rows = []

        sort = request.GET.get("sort", "date")
        if sort not in SORT_FIELDS:
            sort = "date"
        direction = "asc" if request.GET.get("direction") == "asc" else "desc"
        sort_key = SORT_FIELDS[sort]
        present_rows = [row for row in rows if row.get(sort_key) is not None]
        missing_sort_rows = [row for row in rows if row.get(sort_key) is None]
        present_rows.sort(key=lambda row: row.get(sort_key), reverse=(direction == "desc"))
        rows = present_rows + missing_sort_rows

        period_days = max(1, (end_date - start_date).days + 1)
        previous_start = start_date - timedelta(days=period_days)
        previous_end = start_date - timedelta(days=1)
        requested_export = (request.GET.get("export") or request.GET.get("type") or "").lower()
        include_analytics = include_full_analytics or requested_export in {
            "csv_bundle", "xlsx", "json", "pdf", "docx", "png", "svg",
        }
        previous_scope = dict(
            plant=plant,
            start_date=previous_start,
            end_date=previous_end,
            water_type=water_type,
            parameter_name=parameter_name,
            phase_id=phase_id,
            search=search,
        )
        previous_rows = cached_result_rows(**previous_scope) if plant and include_analytics else []

        summaries = parameter_summaries(rows)
        daily_summaries = daily_parameter_summaries(rows)
        source_sample_count = count_analysis_records(
            plant=plant,
            start_date=start_date,
            end_date=end_date,
            water_type=water_type,
            parameter_name=parameter_name,
            phase_id=phase_id,
            search=search,
        ) if plant and not search else None
        samples = sample_summary(rows, samples_collected=source_sample_count)
        analytics = None
        if include_analytics and plant:
            analytics_key = f"qcr:analytics:v2:{plant.pk}:{water_type}:{start_date}:{end_date}:{parameter_name}:{focus_parameter}:{group_by}:{search}"
            analytics = cache.get(analytics_key)
            if analytics is None:
                analytics = build_analytics(rows, previous_rows, focus_parameter, group_by)
                if analytics is not None:
                    analytics["multivariate"] = multivariate_analysis(rows)
                    cache.set(analytics_key, analytics, timeout=600)
        elif include_analytics:
            analytics = build_analytics(rows, previous_rows, focus_parameter, group_by)
            if analytics is not None:
                analytics["multivariate"] = multivariate_analysis(rows)
        total_results = len(rows)
        evaluated = sum(row["compliant"] is not None for row in rows)
        within_limits = sum(row["compliant"] is True for row in rows)
        out_of_range = sum(row["compliant"] is False for row in rows)
        missing = sum(row["status"] == "Missing" for row in rows)
        pending = sum(row["status"] in {"Pending entry", "Verify historical value"} for row in rows)
        not_evaluated = total_results - evaluated - missing - pending
        compliance_rate = round(within_limits / evaluated * 100, 1) if evaluated else None

        return {
            "plant": plant,
            "available_plants": available_plants,
            "start_date": start_date,
            "end_date": end_date,
            "period": period,
            "period_choices": PERIOD_CHOICES,
            "group_by": group_by,
            "group_choices": GROUP_CHOICES,
            "water_type": water_type,
            "water_types": available_water_types,
            "parameters": parameters,
            "parameter": parameter_name,
            "focus_parameter": (analytics or {}).get("focus_parameter", focus_parameter),
            "phases": phases,
            "phase": phase_id,
            "search": search,
            "sort": sort,
            "direction": direction,
            "rows": rows,
            "previous_rows": previous_rows,
            "summaries": summaries,
            "daily_summaries": daily_summaries,
            "sample_summary": samples,
            "analytics": analytics,
            "total_results": total_results,
            "within_limits": within_limits,
            "out_of_range": out_of_range,
            "missing": missing,
            "pending": pending,
            "not_evaluated": not_evaluated,
            "compliance_rate": compliance_rate,
            "analyzed_at": timezone.localtime(),
            "report_metadata": {
                "schema_version": REPORT_SCHEMA_VERSION,
                "generated_at": timezone.localtime().isoformat(),
                "generated_by": request.user.get_username(),
                "treatment_plant_id": plant.pk if plant else None,
                "treatment_plant": plant.name if plant else "",
                "water_type": water_type,
                "date_start": start_date.isoformat(),
                "date_end": end_date.isoformat(),
                "filters": {
                    "parameter": parameter_name,
                    "focus": focus_parameter,
                    "phase_id": phase_id,
                    "search": search,
                    "group_by": group_by,
                },
                "limit_values_are_exported_per_result": True,
                "methodology": "Exploratory descriptive and comparative analysis; inspect the accompanying method notes before operational decisions.",
            },
            "export_formats": (
                ("csv", "Results CSV", "fa-solid fa-file-csv"),
                ("csv_bundle", "CSV bundle", "fa-solid fa-file-zipper"),
                ("xlsx", "Excel", "fa-solid fa-file-excel"),
                ("json", "JSON", "fa-solid fa-code"),
                ("docx", "Word", "fa-solid fa-file-word"),
                ("pdf", "PDF", "fa-solid fa-file-pdf"),
                ("png", "PNG image", "fa-regular fa-image"),
                ("svg", "SVG image", "fa-solid fa-bezier-curve"),
            ),
        }

    def template_context(self, data, page_kind="reports"):
        paginator = Paginator(data["rows"], 60)
        page_obj = paginator.get_page(self.request.GET.get("page", 1))
        daily_paginator = Paginator(data["daily_summaries"], 60)
        daily_page_obj = daily_paginator.get_page(self.request.GET.get("daily_page", 1))
        context = {key: value for key, value in data.items() if key not in {"rows", "previous_rows"}}
        context.update({
            "page_obj": page_obj,
            "page_rows": page_obj.object_list,
            "daily_page_obj": daily_page_obj,
            "daily_rows": daily_page_obj.object_list,
            "page_kind": page_kind,
            "is_reports_page": page_kind == "reports",
            "is_analytics_page": page_kind == "analytics",
            "is_visualizations_page": page_kind == "visualizations",
            "is_qca_insights_page": True,
            "dashboard_url": reverse(f"qc:{page_kind}"),
            "filter_action": reverse(f"qc:{page_kind}"),
            "water_type_label": dict(data["water_types"]).get(data["water_type"], data["water_type"].replace("_", " ").title()),
        })
        query_scope = self.request.GET.copy()
        query_scope.pop("export", None)
        query_scope.pop("type", None)
        water_type_urls = []
        for key, label in data["water_types"]:
            params = query_scope.copy()
            params["water_type"] = key
            params["page"] = "1"
            params["daily_page"] = "1"
            water_type_urls.append((key, label, f"{reverse(f'qc:{page_kind}')}?{params.urlencode()}"))
        context["water_type_urls"] = water_type_urls
        return context


class QCInsightsDashboardView(QCReportsBase):
    page_kind = "analytics"
    template_name = "qc/reports/analytics_dashboard.html"

    def get(self, request, *args, **kwargs):
        data = self.get_report_data(request, include_full_analytics=True)
        return render(request, self.template_name, self.template_context(data, self.page_kind))

    def export_report(self, export_type, data):
        # Reuse the common advanced report exporter implemented below.
        return QCVisualizationDashboardView.export_report(self, export_type, data)


class QCVisualizationDashboardView(QCInsightsDashboardView):
    page_kind = "visualizations"
    template_name = "qc/reports/visualizations_dashboard.html"

    def export_report(self, export_type, data):
        supported = {"csv", "csv_bundle", "xlsx", "json", "docx", "pdf", "png", "svg"}
        if export_type not in supported:
            return HttpResponse("Unsupported export format.", status=400, content_type="text/plain")
        slug = data["water_type"]
        filename = f"qc-{slug}-{data['start_date']}-{data['end_date']}"
        rows = data["rows"]
        fields, table = rows_for_export(rows)
        def csv_bytes(headers, records):
            buffer = io.StringIO(newline="")
            writer = csv.writer(buffer)
            writer.writerow(headers)
            for record in records:
                writer.writerow([record.get(key, "") if record.get(key) is not None else "" for key in headers])
            return buffer.getvalue().encode("utf-8-sig")

        if export_type == "csv":
            return self.download(
                csv_bytes([key for key, _ in fields], rows),
                f"{filename}-results.csv",
                "text/csv; charset=utf-8",
            )

        if export_type == "csv_bundle":
            analytics = data["analytics"] or {}
            parameter_fields = (
                "parameter", "unit", "n", "result_count", "samples_collected", "samples_complying",
                "sample_compliance_rate", "missing", "pending", "not_evaluated", "mean", "median", "std",
                "cv", "p05", "q1", "q3", "p95", "min", "max", "compliant", "out_of_range", "compliance_rate",
            )
            daily_fields = tuple(data["daily_summaries"][0].keys()) if data["daily_summaries"] else (
                "date", "parameter", "unit", "samples_collected", "samples_complying", "result_count", "measured_count",
            )
            metric_rows = []
            for metric in analytics.get("metrics", []):
                metric_rows.append({"section": "metric", "name": metric.get("label"), "value": metric.get("value"), "detail": metric.get("detail")})
            trend = analytics.get("trend", {})
            for key, value in trend.items():
                metric_rows.append({"section": "trend", "name": key, "value": value})
            for item in analytics.get("diagnostics", []):
                metric_rows.append({"section": "diagnostic", "name": item.get("name"), "value": item.get("statistic"), "p_value": item.get("p_value"), "detail": item.get("interpretation")})
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("results.csv", csv_bytes([key for key, _ in fields], rows))
                archive.writestr("parameter_summary.csv", csv_bytes(list(parameter_fields), data["summaries"]))
                archive.writestr("daily_summary.csv", csv_bytes(list(daily_fields), data["daily_summaries"]))
                archive.writestr("analytics.csv", csv_bytes(["section", "name", "value", "p_value", "detail"], metric_rows))
                archive.writestr("metadata.json", json.dumps(json_safe(data["report_metadata"]), ensure_ascii=False, indent=2))
            return self.download(zip_buffer.getvalue(), f"{filename}-report-bundle.zip", "application/zip")

        if export_type == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            summary = workbook.active
            summary.title = "Executive summary"
            summary.append(["QCA advanced report", ""])
            summary.append(["Treatment plant", data["plant"].name if data["plant"] else "No plant assigned"])
            summary.append(["Period", f"{data['start_date']} to {data['end_date']}"])
            summary.append(["Water type", data["water_type"].replace("_", " ").title()])
            summary.append(["Results", data["total_results"]])
            for key, label in (("samples_collected", "Samples collected"), ("samples_complying", "Samples complying"), ("samples_out_of_range", "Samples out of range"), ("samples_evaluated", "Samples fully evaluated")):
                summary.append([label, data["sample_summary"][key]])
            summary.append(["Sample compliance rate", f"{data['sample_summary']['sample_compliance_rate']}%" if data["sample_summary"]["sample_compliance_rate"] is not None else "—"])
            summary.append(["Within limits", data["within_limits"]])
            summary.append(["Out of range", data["out_of_range"]])
            summary.append(["Missing", data["missing"]])
            summary.append(["Pending / historical values to verify", data["pending"]])
            summary.append(["Not evaluated (no limits configured)", data["not_evaluated"]])
            summary.append(["Compliance rate", f"{data['compliance_rate']}%" if data["compliance_rate"] is not None else "—"])
            parameter_sheet = workbook.create_sheet("Parameter statistics")
            stat_headers = ["Parameter", "Unit", "N", "Result cells", "Samples collected", "Samples complying", "Sample compliance %", "Missing", "Pending", "Not evaluated", "Mean", "Median", "SD", "CV %", "P05", "P95", "Min", "Max", "Within limits", "Out of range", "Compliance %", "Mean CI 95% lower", "Mean CI 95% upper"]
            parameter_sheet.append(stat_headers)
            for item in data["summaries"]:
                parameter_sheet.append([item.get(key) for key in ("parameter", "unit", "n", "result_count", "samples_collected", "samples_complying", "sample_compliance_rate", "missing", "pending", "not_evaluated", "mean", "median", "std", "cv", "p05", "p95", "min", "max", "compliant", "out_of_range", "compliance_rate", "ci_low", "ci_high")])
            daily_sheet = workbook.create_sheet("Daily analytics")
            daily_sheet.append(["Date", "Parameter", "Unit", "Analyses in scope", "Samples complying", "Samples out of range", "Sample compliance %", "Result count", "Measured", "Min", "Max", "Mean", "Median", "SD", "Within limits", "Out of range", "Missing", "Pending", "No limits", "Result compliance %"])
            for item in data["daily_summaries"]:
                daily_sheet.append([item.get(key) for key in ("date", "parameter", "unit", "samples_collected", "samples_complying", "samples_out_of_range", "sample_compliance_rate", "result_count", "measured_count", "minimum", "maximum", "mean", "median", "std", "within_limits", "out_of_range", "missing", "pending", "no_limits", "result_compliance_rate")])
            sheet = workbook.create_sheet("Detailed results")
            for row in table:
                sheet.append(row)
            analytics = data["analytics"] or {}
            analysis_sheet = workbook.create_sheet("Advanced analytics")
            analysis_sheet.append(["Parameter analytics", analytics.get("focus_parameter", "—")])
            analysis_sheet.append(["Unit", analytics.get("unit", "")])
            trend = analytics.get("trend", {})
            analysis_sheet.append(["OLS slope per day", trend.get("slope_per_day")])
            analysis_sheet.append(["Robust Theil–Sen slope per day", trend.get("robust_slope")])
            analysis_sheet.append(["Daily median trend HAC p-value", trend.get("p_value")])
            analysis_sheet.append(["Trend R squared", trend.get("r_squared")])
            analysis_sheet.append(["Trend 95% CI lower", trend.get("ci_low")])
            analysis_sheet.append(["Trend 95% CI upper", trend.get("ci_high")])
            forecast_validation = analytics.get("forecast_validation", {})
            analysis_sheet.append(["Forecast rolling-origin validation points", forecast_validation.get("points")])
            analysis_sheet.append(["Forecast validation MAE", forecast_validation.get("mae")])
            analysis_sheet.append(["Forecast validation RMSE", forecast_validation.get("rmse")])
            analysis_sheet.append(["Forecast interval method", forecast_validation.get("interval_method")])
            analysis_sheet.append([])
            analysis_sheet.append(["Period", "n", "Mean", "Median", "Within limits", "Out of range", "Compliance %"])
            for item in analytics.get("by_group", []):
                analysis_sheet.append([item.get("label"), item.get("n"), item.get("mean"), item.get("median"), item.get("within"), item.get("out"), item.get("compliance_rate")])
            if analytics.get("stage_comparison"):
                analysis_sheet.append([])
                analysis_sheet.append(["Water type / phase", "N", "Mean", "Median", "SD", "95% CI lower", "95% CI upper", "Compliance %"])
                for item in analytics["stage_comparison"]:
                    analysis_sheet.append([item.get("label"), item.get("n"), item.get("mean"), item.get("median"), item.get("std"), item.get("ci_low"), item.get("ci_high"), item.get("compliance_rate")])
            analysis_sheet.append([])
            analysis_sheet.append(["Model diagnostic", "Statistic", "p-value", "Interpretation"])
            for item in analytics.get("diagnostics", []):
                analysis_sheet.append([item.get("name"), item.get("statistic"), item.get("p_value"), item.get("interpretation")])
            if analytics.get("stage_test"):
                analysis_sheet.append([])
                analysis_sheet.append([analytics["stage_test"].get("name"), analytics["stage_test"].get("statistic"), analytics["stage_test"].get("p_value"), analytics["stage_test"].get("interpretation")])
            if analytics.get("comparisons"):
                analysis_sheet.append([])
                analysis_sheet.append(["Previous-period comparison", "Value"])
                for key, label in (("difference", "Mean difference"), ("change_percent", "Mean change %"), ("welch_p_value", "Welch p-value"), ("mann_whitney_p_value", "Mann–Whitney p-value"), ("rank_biserial_effect", "Rank-biserial effect"), ("cohens_d", "Cohen’s d")):
                    analysis_sheet.append([label, analytics["comparisons"][1].get(key)])
            if analytics.get("forecast"):
                analysis_sheet.append([])
                analysis_sheet.append(["Damped trend forecast", "Mean", "Approx. 95% lower", "Approx. 95% upper"])
                for item in analytics["forecast"]:
                    analysis_sheet.append([item.get("period"), item.get("mean"), item.get("low"), item.get("high")])
            anomalies = analytics.get("anomalies", [])
            if anomalies:
                analysis_sheet.append([])
                analysis_sheet.append(["Anomaly candidates", "Sample", "Date", "Time", "Parameter", "Result", "Unit", "Phase"])
            for item in anomalies:
                analysis_sheet.append(["Isolation Forest", item.get("reference"), item.get("date"), item.get("time"), item.get("parameter"), item.get("value"), item.get("unit"), item.get("phase")])
            multivariate = analytics.get("multivariate", {})
            if multivariate.get("correlations"):
                analysis_sheet.append([])
                analysis_sheet.append(["Cross-parameter association", "Parameter A", "Parameter B", "n", "Pearson r", "Pearson p", "Pearson FDR q", "Spearman rho", "Spearman p", "Spearman FDR q"])
                rank_lookup = {(item["left"], item["right"]): item for item in multivariate.get("spearman_correlations", [])}
                for item in multivariate["correlations"]:
                    rank = rank_lookup.get((item["left"], item["right"]), {})
                    analysis_sheet.append(["Pairwise complete samples", item["left"], item["right"], item["n"], item["r"], item["p_value"], item.get("fdr_q_value"), rank.get("rho"), rank.get("p_value"), rank.get("fdr_q_value")])
            pca = multivariate.get("pca", {})
            if pca.get("available"):
                analysis_sheet.append([])
                analysis_sheet.append(["PCA explained variance", "Proportion", "Samples", pca.get("samples_used")])
                for index, value in enumerate(pca.get("explained_variance", []), start=1):
                    analysis_sheet.append([f"PC{index}", value])
                analysis_sheet.append([])
                analysis_sheet.append(["PCA loading", "PC1", "PC2"])
                for item in pca.get("loadings", []):
                    analysis_sheet.append([item.get("parameter"), item.get("pc1"), item.get("pc2")])
            for ws in workbook.worksheets:
                ws.freeze_panes = "A2"
                if ws.title in {"Parameter statistics", "Daily analytics", "Detailed results"}:
                    ws.auto_filter.ref = ws.dimensions
                for row_index, row in enumerate(ws.iter_rows(), start=1):
                    first = row[0].value if row else None
                    if (ws.title in {"Detailed results", "Daily analytics", "Parameter statistics"} and row_index == 1) or first in {"QCA advanced report", "Parameter statistics", "Parameter analytics", "Period", "Water type / phase", "Model diagnostic", "Previous-period comparison", "Damped trend forecast", "Anomaly candidates"}:
                        for cell in row:
                            cell.font = Font(color="FFFFFF", bold=True)
                            cell.fill = PatternFill("solid", fgColor="12304A")
                for col in ws.columns:
                    width = min(36, max(12, max(len(str(cell.value or "")) for cell in list(col)[:250]) + 2))
                    ws.column_dimensions[get_column_letter(col[0].column)].width = width
                    for cell in list(col)[1:]:
                        cell.alignment = Alignment(vertical="top")
            metadata_sheet = workbook.create_sheet("Report metadata")
            metadata_sheet.append(["Field", "Value"])
            for key, value in data["report_metadata"].items():
                metadata_sheet.append([key, json.dumps(json_safe(value), ensure_ascii=False) if isinstance(value, (dict, list, tuple)) else str(value)])
            metadata_sheet.freeze_panes = "A2"
            output = io.BytesIO()
            workbook.save(output)
            return self.download(output.getvalue(), f"{filename}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        if export_type == "json":
            payload = {
                "metadata": data["report_metadata"],
                "summary": {key: data[key] for key in ("total_results", "within_limits", "out_of_range", "missing", "pending", "not_evaluated", "compliance_rate")},
                "sample_summary": data["sample_summary"],
                "parameter_statistics": data["summaries"],
                "daily_analytics": data["daily_summaries"],
                "analytics": data["analytics"],
                "results": data["rows"],
            }
            body = json.dumps(json_safe(payload), ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8")
            return self.download(body, f"{filename}.json", "application/json; charset=utf-8")

        if export_type in {"png", "svg"}:
            analytics = data["analytics"] or build_analytics(data["rows"], data["previous_rows"], data["focus_parameter"], data["group_by"])
            output = image_report(data["rows"], analytics, export_type)
            mime = "image/png" if export_type == "png" else "image/svg+xml"
            return self.download(output.getvalue(), f"{filename}.{export_type}", mime)

        if export_type == "docx":
            from docx import Document
            from docx.shared import Inches

            analytics = data["analytics"] or build_analytics(data["rows"], data["previous_rows"], data["focus_parameter"], data["group_by"])
            document = Document()
            document.add_heading("QCA Advanced Water Quality Report", 0)
            document.add_paragraph(f"Plant: {data['plant'].name if data['plant'] else '—'}")
            document.add_paragraph(f"Period: {data['start_date']} through {data['end_date']} · Water type: {data['water_type'].replace('_', ' ').title()}")
            document.add_heading("Executive summary", level=1)
            sample_stats = data["sample_summary"]
            document.add_paragraph(f"{sample_stats['samples_collected']:,} samples collected; {sample_stats['samples_complying']:,} fully complying; {sample_stats['samples_out_of_range']:,} with at least one out-of-range result; {sample_stats['samples_not_fully_evaluated']:,} not fully evaluated. Result cells: {data['total_results']:,}; {data['within_limits']:,} within limits; {data['out_of_range']:,} out of range; {data['missing']:,} missing; {data['pending']:,} pending or unverified.")
            if analytics.get("trend", {}).get("slope_per_day") is not None:
                trend = analytics["trend"]
                document.add_paragraph(f"Focused parameter: {analytics['focus_parameter']}. Daily median trend: {trend['slope_per_day']:.5g} {analytics.get('unit','')} per day (HAC p={trend['p_value']:.4g}, R²={trend['r_squared']:.3f}); Theil–Sen slope {trend['robust_slope']:.5g} per day.")
            document.add_heading("Parameter summary statistics", level=1)
            summary_table = document.add_table(rows=1, cols=13)
            summary_table.style = "Light Shading Accent 1"
            for cell, label in zip(summary_table.rows[0].cells, ("Parameter", "Unit", "N", "Samples", "Complying samples", "Min", "Max", "Missing", "No limits", "Mean", "Median", "Compliance", "Out of range")):
                cell.text = label
            for item in data["summaries"]:
                cells = summary_table.add_row().cells
                values = (item["parameter"], item["unit"], item["n"], item["samples_collected"], item["samples_complying"], item["min"], item["max"], item["missing"], item["not_evaluated"], item["mean"], item["median"], f"{item['compliance_rate']}%" if item["compliance_rate"] is not None else "—", item["out_of_range"])
                for cell, value in zip(cells, values):
                    cell.text = str(value if value is not None else "—")
            document.add_heading("Daily descriptive analytics", level=1)
            daily_table = document.add_table(rows=1, cols=11)
            daily_table.style = "Light Shading Accent 1"
            for cell, label in zip(daily_table.rows[0].cells, ("Date", "Parameter", "Unit", "Samples", "Complying", "Min", "Max", "Mean", "Median", "Out of range", "Compliance %")):
                cell.text = label
            for item in data["daily_summaries"]:
                cells = daily_table.add_row().cells
                values = (item["date"], item["parameter"], item["unit"], item["samples_collected"], item["samples_complying"], item["minimum"], item["maximum"], item["mean"], item["median"], item["out_of_range"], item["result_compliance_rate"])
                for cell, value in zip(cells, values):
                    cell.text = str(value if value is not None else "—")
            document.add_heading("Advanced statistical analysis", level=1)
            document.add_paragraph(f"Analytics focus: {analytics.get('focus_parameter', '—')} {analytics.get('unit', '')}. Grouping interval: {data['group_by']}.")
            trend = analytics.get("trend", {})
            if trend.get("slope_per_day") is not None:
                document.add_paragraph(f"Daily median trend slope {trend['slope_per_day']:.6g} per day (HAC robust 95% CI {trend['ci_low']:.6g} to {trend['ci_high']:.6g}; p={trend['p_value']:.5g}; R²={trend['r_squared']:.4f}). Robust Theil–Sen slope: {trend.get('robust_slope'):.6g} per day.")
            comparisons = analytics.get("comparisons", [])
            if len(comparisons) > 1:
                previous = comparisons[1]
                document.add_paragraph(f"Previous-window comparison: mean difference {previous.get('difference')}; Welch p={previous.get('welch_p_value')}; Mann–Whitney p={previous.get('mann_whitney_p_value')}; rank-biserial effect={previous.get('rank_biserial_effect')}; Cohen’s d={previous.get('cohens_d')}.")
            if analytics.get("stage_test"):
                stage_test = analytics["stage_test"]
                document.add_paragraph(f"{stage_test['name']}: statistic {stage_test['statistic']:.5g}, p={stage_test['p_value']:.5g}. {stage_test['interpretation']}.")
            if analytics.get("stage_comparison"):
                document.add_heading("Water stage / phase comparison", level=2)
                stage_table = document.add_table(rows=1, cols=7)
                stage_table.style = "Light Shading Accent 1"
                for cell, label in zip(stage_table.rows[0].cells, ("Water type / phase", "N", "Mean", "Median", "SD", "Mean CI 95%", "Compliance")):
                    cell.text = label
                for item in analytics["stage_comparison"]:
                    cells = stage_table.add_row().cells
                    values = (item["label"], item["n"], item["mean"], item["median"], item["std"], f"{item['ci_low']} to {item['ci_high']}", f"{item['compliance_rate']}%" if item["compliance_rate"] is not None else "Not evaluated")
                    for cell, value in zip(cells, values):
                        cell.text = str(value if value is not None else "—")
            multivariate = analytics.get("multivariate", {})
            if multivariate.get("correlations"):
                document.add_heading("Cross-parameter association", level=2)
                document.add_paragraph("Pairwise complete sample associations; Benjamini–Hochberg q-values adjust for multiple comparisons.")
                association_table = document.add_table(rows=1, cols=8)
                association_table.style = "Light Shading Accent 1"
                for cell, label in zip(association_table.rows[0].cells, ("Parameter A", "Parameter B", "n", "Pearson r", "Pearson q", "Spearman ρ", "Spearman q", "Pearson p")):
                    cell.text = label
                rank_lookup = {(item["left"], item["right"]): item for item in multivariate.get("spearman_correlations", [])}
                for item in multivariate["correlations"]:
                    rank = rank_lookup.get((item["left"], item["right"]), {})
                    cells = association_table.add_row().cells
                    values = (item["left"], item["right"], item["n"], item["r"], item.get("fdr_q_value"), rank.get("rho"), rank.get("fdr_q_value"), item.get("p_value"))
                    for cell, value in zip(cells, values):
                        cell.text = str(value if value is not None else "—")
            pca = multivariate.get("pca", {})
            if pca.get("available"):
                document.add_heading("Principal component analysis", level=2)
                variance_text = ", ".join(f"PC{index}: {value:.1%}" for index, value in enumerate(pca.get("explained_variance", []), start=1))
                document.add_paragraph(f"{pca.get('samples_used')} sample profiles. Explained variance: {variance_text}. {pca.get('reason')}")
                loading_table = document.add_table(rows=1, cols=3)
                loading_table.style = "Light Shading Accent 1"
                for cell, label in zip(loading_table.rows[0].cells, ("Parameter", "PC1 loading", "PC2 loading")):
                    cell.text = label
                for item in pca.get("loadings", []):
                    cells = loading_table.add_row().cells
                    for cell, value in zip(cells, (item.get("parameter"), item.get("pc1"), item.get("pc2"))):
                        cell.text = str(value if value is not None else "—")
            diagnostic_table = document.add_table(rows=1, cols=3)
            diagnostic_table.style = "Light Shading Accent 1"
            for cell, label in zip(diagnostic_table.rows[0].cells, ("Diagnostic", "Statistic / p-value", "Interpretation")):
                cell.text = label
            for item in analytics.get("diagnostics", []):
                cells = diagnostic_table.add_row().cells
                statistic = f"{item.get('statistic')} / p={item.get('p_value')}" if item.get("p_value") is not None else str(item.get("statistic"))
                for cell, value in zip(cells, (item.get("name"), statistic, item.get("interpretation"))):
                    cell.text = str(value or "—")
            if data["rows"]:
                figure = image_report(data["rows"], analytics, "png")
                document.add_picture(figure, width=Inches(6.3))
            document.add_heading("Detailed sample results", level=1)
            result_table = document.add_table(rows=1, cols=9)
            result_table.style = "Light Shading Accent 1"
            labels = ("Ref", "Date", "Time", "Type", "Phase", "Parameter", "Result", "Limits", "Status")
            for cell, label in zip(result_table.rows[0].cells, labels):
                cell.text = label
            for row in data["rows"]:
                cells = result_table.add_row().cells
                values = (row["reference"], row["date"], row["time"], row["water_type_label"], row["phase"], row["parameter"], f"{row['value_text']} {row['unit']}".strip(), f"{row['minimum_text'] or '—'} to {row['maximum_text'] or '—'}", row["status"])
                for cell, value in zip(cells, values):
                    cell.text = str(value)
            document.add_paragraph("Statistical flags are screening aids. Confirm anomalies and compliance interpretation against validated laboratory methods and applicable standards.")
            output = io.BytesIO()
            document.save(output)
            return self.download(output.getvalue(), f"{filename}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        if export_type == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import landscape, legal
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import Image, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle
            from matplotlib import font_manager
            from xml.sax.saxutils import escape

            analytics = data["analytics"] or build_analytics(data["rows"], data["previous_rows"], data["focus_parameter"], data["group_by"])
            output = io.BytesIO()
            doc = SimpleDocTemplate(output, pagesize=landscape(legal), rightMargin=0.35*inch, leftMargin=0.35*inch, topMargin=0.45*inch, bottomMargin=0.45*inch, title="QCA Advanced Water Quality Report")
            styles = getSampleStyleSheet()
            report_font = font_manager.findfont("DejaVu Sans")
            if "QCAReportFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("QCAReportFont", report_font))
            for style_name in ("Title", "BodyText", "Heading2"):
                styles[style_name].fontName = "QCAReportFont"
            styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], textColor=colors.HexColor("#12304a"), alignment=TA_CENTER, spaceAfter=8))
            styles.add(ParagraphStyle(name="SmallCell", parent=styles["BodyText"], fontName="QCAReportFont", fontSize=6.5, leading=8))
            plant_name = escape(data["plant"].name if data["plant"] else "—")
            story = [Paragraph("QCA Advanced Water Quality Report", styles["ReportTitle"]), Paragraph(f"<b>Plant:</b> {plant_name} &nbsp;&nbsp; <b>Period:</b> {data['start_date']} – {data['end_date']} &nbsp;&nbsp; <b>Water type:</b> {escape(data['water_type'].replace('_', ' ').title())}", styles["BodyText"]), Spacer(1, 8)]
            sample_stats = data["sample_summary"]
            story.append(Paragraph(f"<b>{sample_stats['samples_collected']:,}</b> samples &nbsp; · &nbsp; <b>{sample_stats['samples_complying']:,}</b> fully complying &nbsp; · &nbsp; <b>{sample_stats['samples_out_of_range']:,}</b> out of range &nbsp; · &nbsp; <b>{data['total_results']:,}</b> results &nbsp; · &nbsp; <b>{data['pending']:,}</b> pending / unverified &nbsp; · &nbsp; <b>{data['compliance_rate'] if data['compliance_rate'] is not None else '—'}%</b> result compliance", styles["BodyText"]))
            story.append(Spacer(1, 10))
            if data["rows"]:
                figure = image_report(data["rows"], analytics, "png")
                story.append(Image(figure, width=9.2*inch, height=6.1*inch, kind="proportional"))
                story.append(PageBreak())
            story.append(Paragraph("Parameter summary statistics", styles["Heading2"]))
            summary_rows = [["Parameter", "Unit", "N", "Samples", "Complying samples", "Sample compliance %", "Observed min", "Observed max", "Missing", "No limits", "Mean", "Median", "SD", "CV %", "P05", "P95", "Compliance", "Out of range"]]
            for item in data["summaries"]:
                summary_rows.append([item.get("parameter"), item.get("unit"), item.get("n"), item.get("samples_collected"), item.get("samples_complying"), item.get("sample_compliance_rate"), item.get("min"), item.get("max"), item.get("missing"), item.get("not_evaluated"), item.get("mean"), item.get("median"), item.get("std"), item.get("cv"), item.get("p05"), item.get("p95"), item.get("compliance_rate"), item.get("out_of_range")])
            story.append(LongTable(summary_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12304a")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 7), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f1f5f9")]), ("VALIGN", (0,0), (-1,-1), "TOP")])))
            story.append(Spacer(1, 8))
            story.append(Paragraph("Daily descriptive analytics", styles["Heading2"]))
            daily_rows = [["Date", "Parameter", "Unit", "Samples", "Complying", "Min", "Max", "Mean", "Median", "Out", "Compliance %"]]
            daily_rows.extend([[item.get("date"), item.get("parameter"), item.get("unit"), item.get("samples_collected"), item.get("samples_complying"), item.get("minimum"), item.get("maximum"), item.get("mean"), item.get("median"), item.get("out_of_range"), item.get("result_compliance_rate")] for item in data["daily_summaries"]])
            story.append(LongTable(daily_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#334155")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 7), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f1f5f9")]), ("VALIGN", (0,0), (-1,-1), "TOP")])))
            story.append(Spacer(1, 10))
            story.append(Paragraph("Advanced statistical analysis", styles["Heading2"]))
            trend = analytics.get("trend", {})
            if trend.get("slope_per_day") is not None:
                story.append(Paragraph(f"<b>Focus:</b> {escape(str(analytics.get('focus_parameter', '—')))} {escape(str(analytics.get('unit', '')))}. <b>Daily median slope:</b> {trend['slope_per_day']:.6g}/day; <b>HAC 95% CI:</b> {trend['ci_low']:.6g} to {trend['ci_high']:.6g}; <b>p:</b> {trend['p_value']:.5g}; <b>R²:</b> {trend['r_squared']:.4f}; <b>Theil–Sen:</b> {trend.get('robust_slope'):.6g}/day.", styles["BodyText"]))
            comparisons = analytics.get("comparisons", [])
            if len(comparisons) > 1:
                previous = comparisons[1]
                story.append(Paragraph(f"<b>Previous-window difference:</b> {previous.get('difference')} {escape(str(analytics.get('unit', '')))}; Welch p={previous.get('welch_p_value')}; Mann–Whitney p={previous.get('mann_whitney_p_value')}; rank-biserial effect={previous.get('rank_biserial_effect')}; Cohen’s d={previous.get('cohens_d')}.", styles["BodyText"]))
            if analytics.get("stage_test"):
                stage_test = analytics["stage_test"]
                story.append(Paragraph(f"<b>{escape(stage_test['name'])}:</b> H={stage_test['statistic']:.5g}, p={stage_test['p_value']:.5g}. {escape(stage_test['interpretation'])}.", styles["BodyText"]))
            if analytics.get("stage_comparison"):
                stage_rows = [["Water type / phase", "N", "Mean", "Median", "SD", "95% CI lower", "95% CI upper", "Compliance %"]]
                for item in analytics["stage_comparison"]:
                    stage_rows.append([item.get("label"), item.get("n"), item.get("mean"), item.get("median"), item.get("std"), item.get("ci_low"), item.get("ci_high"), item.get("compliance_rate")])
                story.append(Spacer(1, 5))
                story.append(LongTable(stage_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#334155")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 7), ("VALIGN", (0,0), (-1,-1), "TOP")])))
            multivariate = analytics.get("multivariate", {})
            if multivariate.get("correlations"):
                story.append(Spacer(1, 8))
                story.append(Paragraph("Cross-parameter associations · pairwise complete samples; Benjamini–Hochberg q-values adjust for multiple comparisons", styles["Heading2"]))
                rank_lookup = {(item["left"], item["right"]): item for item in multivariate.get("spearman_correlations", [])}
                association_rows = [["Parameter A", "Parameter B", "n", "Pearson r", "Pearson p", "FDR q", "Spearman rho"]]
                for item in multivariate["correlations"]:
                    rank = rank_lookup.get((item["left"], item["right"]), {})
                    association_rows.append([item["left"], item["right"], item["n"], item["r"], item["p_value"], item.get("fdr_q_value"), rank.get("rho")])
                story.append(LongTable(association_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#334155")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 7), ("VALIGN", (0,0), (-1,-1), "TOP")])))
            pca = multivariate.get("pca", {})
            if pca.get("available"):
                explained = ", ".join(f"PC{index}: {value:.1%}" for index, value in enumerate(pca.get("explained_variance", []), start=1))
                story.append(Spacer(1, 5))
                story.append(Paragraph(f"<b>Principal component analysis:</b> {pca.get('samples_used')} sample profiles · explained variance {escape(explained)}. {escape(pca.get('reason', ''))}", styles["BodyText"]))
            diagnostic_rows = [["Model diagnostic", "Statistic", "p-value", "Interpretation"]]
            for item in analytics.get("diagnostics", []):
                diagnostic_rows.append([item.get("name"), item.get("statistic"), item.get("p_value"), item.get("interpretation")])
            if len(diagnostic_rows) > 1:
                story.append(Spacer(1, 5))
                story.append(LongTable(diagnostic_rows, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#334155")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 7), ("VALIGN", (0,0), (-1,-1), "TOP")])))
            story.append(Spacer(1, 12))
            story.append(Paragraph("Detailed sample results", styles["Heading2"]))
            report_rows = [[Paragraph(cell, styles["SmallCell"]) for cell in table[0]]]
            for row in table[1:]:
                report_rows.append([Paragraph(str(cell).replace("&", "&amp;").replace("<", "&lt;"), styles["SmallCell"]) for cell in row])
            widths = [0.55*inch, 0.9*inch, 0.85*inch, 0.65*inch, 0.48*inch, 0.55*inch, 0.58*inch, 0.82*inch, 0.58*inch, 0.62*inch, 0.7*inch, 0.38*inch, 0.68*inch, 0.65*inch, 0.65*inch, 0.8*inch, 1.1*inch]
            story.append(LongTable(report_rows, colWidths=widths, repeatRows=1, style=TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12304a")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#cbd5e1")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f1f5f9")]), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3)])))
            doc.build(story)
            return self.download(output.getvalue(), f"{filename}.pdf", "application/pdf")

        return HttpResponse("Unsupported export format.", status=400, content_type="text/plain")

    @staticmethod
    def download(content, filename, content_type):
        response = HttpResponse(content, content_type=content_type)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        return response


class QCReportsView(QCVisualizationDashboardView):
    """Per-water-type evidence reports with multi-format export support."""

    page_kind = "reports"
    template_name = "qc/reports/dashboard.html"

    def get(self, request, *args, **kwargs):
        data = self.get_report_data(request)
        export_type = (request.GET.get("export") or "").strip().lower()
        if export_type:
            return self.export_report(export_type, data)
        return render(request, self.template_name, self.template_context(data, self.page_kind))

