from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import timedelta

from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone

from qc.services.advanced_reports import PERIOD_LENGTHS, WATER_TYPES, json_safe
from qc.services.noncompliance import analyze_noncompliance, load_noncompliance_cases
from qc.views.reports import GROUP_CHOICES, PERIOD_CHOICES, QCReportsBase


class QCNonComplianceDashboardView(QCReportsBase):
    """Case dashboard with water-type, report, analytics and visualization tabs."""

    template_name = "qc/noncompliance/dashboard.html"

    def get_report_data(self, request, include_full_analytics=False):
        today = timezone.localdate()
        period = request.GET.get("period", "month")
        if period not in dict(PERIOD_CHOICES):
            period = "month"
        end_date = self.parse_date(request.GET.get("end"), today)
        if period in PERIOD_LENGTHS:
            start_date = end_date - timedelta(days=PERIOD_LENGTHS[period] - 1)
        else:
            start_date = self.parse_date(request.GET.get("start"), end_date - timedelta(days=29))
        if start_date > end_date:
            start_date, end_date = end_date, start_date

        water_types = self.permitted_water_types(request.user)
        permitted = {key for key, _label in water_types}
        water_type = request.GET.get("water_type", "raw")
        if water_type not in permitted:
            water_type = water_types[0][0] if water_types else "raw"
        section = request.GET.get("section", "reports")
        if section not in {"reports", "analytics", "visualizations"}:
            section = "reports"
        group_by = request.GET.get("group_by", "day")
        if group_by not in dict(GROUP_CHOICES):
            group_by = "day"
        search = request.GET.get("search", "").strip()[:120]
        status = request.GET.get("status", "").strip()
        if status not in {"", "open", "under_review", "closed", "derived"}:
            status = ""
        severity = request.GET.get("severity", "").strip()
        if severity not in {"", "Critical", "High", "Moderate", "Low"}:
            severity = ""
        sort = request.GET.get("sort", "date")
        if sort not in {"date", "risk", "deviation", "status", "parameter", "reference"}:
            sort = "date"
        direction = "asc" if request.GET.get("direction") == "asc" else "desc"
        plant, available_plants = self.selected_plant(request)
        cases = load_noncompliance_cases(
            plant=plant, start_date=start_date, end_date=end_date,
            water_type=water_type, status=status, severity=severity,
            search=search, sort=sort, direction=direction,
        )
        analytics = analyze_noncompliance(cases, group_by=group_by, start_date=start_date, end_date=end_date)
        query_scope = request.GET.copy()
        query_scope.pop("export", None)
        query_scope.pop("type", None)

        def section_url(target):
            params = query_scope.copy()
            params["section"] = target
            return f"{reverse('qc:noncompliance')}?{params.urlencode()}"

        def water_type_url(target):
            params = query_scope.copy()
            params["water_type"] = target
            params["page"] = "1"
            return f"{reverse('qc:noncompliance')}?{params.urlencode()}"

        return {
            "plant": plant,
            "available_plants": available_plants,
            "water_types": water_types,
            "water_type": water_type,
            "water_type_label": dict(water_types).get(water_type, water_type.replace("_", " ").title()),
            "start_date": start_date,
            "end_date": end_date,
            "period": period,
            "period_choices": PERIOD_CHOICES,
            "group_by": group_by,
            "group_choices": GROUP_CHOICES,
            "section": section,
            "search": search,
            "status": status,
            "severity": severity,
            "sort": sort,
            "direction": direction,
            "cases": cases,
            "analytics": analytics,
            "analyzed_at": timezone.localtime(),
            "report_metadata": {
                "schema_version": "2.0",
                "generated_at": timezone.localtime().isoformat(),
                "generated_by": request.user.get_username(),
                "plant_id": plant.pk if plant else None,
                "plant": plant.name if plant else "",
                "water_type": water_type,
                "date_start": start_date.isoformat(),
                "date_end": end_date.isoformat(),
                "date_basis": "case recorded_at; filtered-water derived exceedances use analysis timestamp",
                "filters": {"status": status, "severity": severity, "search": search, "group_by": group_by},
                "filtered_water_source": "Derived from recorded results outside configured limits; not a persisted case model.",
                "methodology": analytics.get("method_note", ""),
            },
            "section_urls": {
                "reports": section_url("reports"),
                "analytics": section_url("analytics"),
                "visualizations": section_url("visualizations"),
            },
            "water_type_urls": [(key, label, water_type_url(key)) for key, label in water_types],
            "export_formats": (
                ("csv", "Cases CSV", "fa-solid fa-file-csv"),
                ("csv_bundle", "CSV bundle", "fa-solid fa-file-zipper"),
                ("xlsx", "Excel", "fa-solid fa-file-excel"),
                ("json", "JSON", "fa-solid fa-code"),
                ("docx", "Word", "fa-solid fa-file-word"),
                ("pdf", "PDF", "fa-solid fa-file-pdf"),
                ("png", "PNG image", "fa-regular fa-image"),
                ("svg", "SVG image", "fa-solid fa-bezier-curve"),
            ),
        }

    def get(self, request, *args, **kwargs):
        data = self.get_report_data(request)
        return render(request, self.template_name, self.template_context(data))

    def template_context(self, data, page_kind="noncompliance"):
        paginator = Paginator(data["cases"], 75)
        page_obj = paginator.get_page(self.request.GET.get("page", 1))
        context = {key: value for key, value in data.items() if key != "cases"}
        context.update({
            "cases_page": page_obj,
            "page_cases": page_obj.object_list,
            "dashboard_url": reverse("qc:noncompliance"),
            "filter_action": reverse("qc:noncompliance"),
            "is_noncompliance_page": True,
            "is_qca_insights_page": True,
            "is_reports_page": False,
            "is_analytics_page": False,
            "is_visualizations_page": False,
            "all_cases": data["cases"] if data["section"] == "visualizations" else [],
        })
        return context

    def export_report(self, export_type, data):
        allowed = {"csv", "csv_bundle", "xlsx", "json", "docx", "pdf", "png", "svg"}
        if export_type not in allowed:
            return HttpResponse("Unsupported export format.", status=400, content_type="text/plain")
        filename = f"qca-noncompliance-{data['water_type']}-{data['start_date']}-{data['end_date']}"
        cases = data["cases"]
        analytics = data["analytics"]
        fields = (
            "id", "source", "source_label", "water_type", "date", "sample_date", "time",
            "reference", "phase", "parameter", "unit", "observed_value", "minimum_limit",
            "maximum_limit", "deviation", "risk_score", "risk_band", "status", "status_label", "escalation_level",
            "auto_detected", "resolved_at", "water_filter", "causes", "root_cause", "corrective_action", "preventive_action", "log_history",
        )

        def write_csv(headers, rows):
            buffer = io.StringIO(newline="")
            writer = csv.writer(buffer)
            writer.writerow(headers)
            for row in rows:
                writer.writerow([row.get(field, "") if row.get(field) is not None else "" for field in headers])
            return buffer.getvalue().encode("utf-8-sig")

        if export_type == "csv":
            return self.download(write_csv(fields, cases), f"{filename}-cases.csv", "text/csv; charset=utf-8")
        if export_type == "csv_bundle":
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("cases.csv", write_csv(fields, cases))
                archive.writestr("status_summary.csv", write_csv(("label", "count"), analytics["status_counts"]))
                archive.writestr("risk_summary.csv", write_csv(("label", "count"), analytics["risk_counts"]))
                archive.writestr("parameter_summary.csv", write_csv(("parameter", "cases", "median_risk", "mean_deviation"), analytics["top_parameters"]))
                archive.writestr("period_summary.csv", write_csv(("date", "label", "cases"), analytics["periods"]))
                archive.writestr("aging_summary.csv", write_csv(("label", "count"), analytics["aging_buckets"]))
                archive.writestr("cause_summary.csv", write_csv(("cause", "cases"), analytics["top_causes"]))
                archive.writestr("metadata.json", json.dumps(json_safe(data["report_metadata"]), ensure_ascii=False, indent=2))
            return self.download(zip_buffer.getvalue(), f"{filename}-bundle.zip", "application/zip")
        if export_type == "json":
            body = json.dumps(json_safe({
                "metadata": data["report_metadata"],
                "summary": {key: value for key, value in analytics.items() if key not in {"periods", "status_counts", "risk_counts", "top_parameters"}},
                "periods": analytics["periods"], "status_summary": analytics["status_counts"],
                "risk_summary": analytics["risk_counts"], "parameter_summary": analytics["top_parameters"],
                "cases": cases,
            }), ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8")
            return self.download(body, f"{filename}.json", "application/json; charset=utf-8")
        if export_type == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            summary = workbook.active
            summary.title = "Executive summary"
            summary.append(["Non-compliance report", "Value"])
            summary.append(["Treatment plant", data["plant"].name if data["plant"] else ""])
            summary.append(["Water type", data["water_type_label"]])
            summary.append(["Recorded period", f"{data['start_date']} to {data['end_date']}"])
            for key in ("total", "open", "under_review", "closed", "derived", "critical_high", "risk_mean", "risk_median", "resolution_median_days", "active_age_median_days"):
                summary.append([key.replace("_", " ").title(), analytics.get(key)])
            for title, headers, records in (
                ("Cases", fields, cases),
                ("Period counts", ("date", "label", "cases"), analytics["periods"]),
                ("Status counts", ("label", "count"), analytics["status_counts"]),
                ("Risk bands", ("label", "count"), analytics["risk_counts"]),
                ("Top parameters", ("parameter", "cases", "median_risk", "mean_deviation"), analytics["top_parameters"]),
                ("Active case aging", ("label", "count"), analytics["aging_buckets"]),
                ("Recorded causes", ("cause", "cases"), analytics["top_causes"]),
            ):
                sheet = workbook.create_sheet(title)
                sheet.append(list(headers))
                for row in records:
                    sheet.append([row.get(key) for key in headers])
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
            metadata = workbook.create_sheet("Metadata")
            metadata.append(["Field", "Value"])
            for key, value in data["report_metadata"].items():
                metadata.append([key, json.dumps(json_safe(value), ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)])
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    if row and row[0].row == 1:
                        for cell in row:
                            cell.font = Font(color="FFFFFF", bold=True)
                            cell.fill = PatternFill("solid", fgColor="12304A")
                for column in sheet.columns:
                    width = min(40, max(12, max(len(str(cell.value or "")) for cell in list(column)[:200]) + 2))
                    sheet.column_dimensions[get_column_letter(column[0].column)].width = width
                    for cell in list(column)[1:]:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
            buffer = io.BytesIO()
            workbook.save(buffer)
            return self.download(buffer.getvalue(), f"{filename}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if export_type == "docx":
            from docx import Document

            document = Document()
            document.add_heading(f"{data['water_type_label']} non-compliance report", 0)
            document.add_paragraph(f"Plant: {data['plant'].name if data['plant'] else '—'} · Recorded: {data['start_date']} to {data['end_date']}")
            document.add_paragraph(f"{analytics['total']} cases · {analytics['open']} open · {analytics['under_review']} under review · {analytics['closed']} closed · {analytics['derived']} derived exceedances")
            document.add_paragraph(analytics["method_note"])
            table = document.add_table(rows=1, cols=9)
            table.style = "Light Shading Accent 1"
            for cell, label in zip(table.rows[0].cells, ("Date", "Reference", "Parameter", "Observed", "Limits", "Deviation", "Risk", "Band", "Status")):
                cell.text = label
            for row in cases:
                cells = table.add_row().cells
                values = (row["date"], row["reference"], row["parameter"], row["observed_value"], f"{row['minimum_limit']} – {row['maximum_limit']}", row["deviation"], row["risk_score"], row["risk_band"], row["status_label"])
                for cell, value in zip(cells, values):
                    cell.text = str(value if value is not None else "—")
            buffer = io.BytesIO()
            document.save(buffer)
            return self.download(buffer.getvalue(), f"{filename}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        if export_type == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
            from xml.sax.saxutils import escape

            buffer = io.BytesIO()
            document = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=.35*inch, leftMargin=.35*inch, topMargin=.4*inch, bottomMargin=.4*inch)
            styles = getSampleStyleSheet()
            story = [Paragraph(f"{escape(data['water_type_label'])} non-compliance report", styles["Title"]), Spacer(1, 6)]
            story.append(Paragraph(f"Plant: {escape(data['plant'].name if data['plant'] else '—')} · Recorded: {data['start_date']} to {data['end_date']} · {analytics['total']} cases", styles["BodyText"]))
            story.append(Paragraph(escape(analytics["method_note"]), styles["BodyText"]))
            table_rows = [["Date", "Reference", "Parameter", "Observed", "Limits", "Deviation", "Risk", "Band", "Status"]]
            for row in cases:
                values = [row.get(key) for key in ("date", "reference", "parameter", "observed_value")]
                values.append(f"{row.get('minimum_limit')} – {row.get('maximum_limit')}")
                values.extend(row.get(key) for key in ("deviation", "risk_score", "risk_band", "status_label"))
                table_rows.append([str(value) if value is not None else "—" for value in values])
            story.append(Spacer(1, 8))
            story.append(LongTable(table_rows, repeatRows=1, style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ])))
            document.build(story)
            return self.download(buffer.getvalue(), f"{filename}.pdf", "application/pdf")
        if export_type in {"png", "svg"}:
            import matplotlib
            matplotlib.use("Agg", force=True)
            import matplotlib.pyplot as plt
            import numpy as np
            import pandas as pd

            figure, axes = plt.subplots(2, 2, figsize=(14, 9), constrained_layout=True)
            periods = analytics["periods"][-36:]
            axes[0, 0].plot([row["label"] for row in periods], [row["cases"] for row in periods], color="#0072B2", marker="o", label="Recorded cases")
            if analytics.get("trend"):
                trend_points = list(range(len(analytics["periods"]))) [-len(periods):]
                slope = analytics["trend"]["slope_per_period"]
                intercept = analytics["trend"]["intercept"]
                axes[0, 0].plot([row["label"] for row in periods], [intercept + slope * point for point in trend_points], color="#CC79A7", linestyle="--", label="Theil–Sen trend")
            axes[0, 0].set_title("Cases recorded per period")
            axes[0, 0].tick_params(axis="x", rotation=35)
            axes[0, 0].legend(frameon=False, fontsize=8)
            risk = analytics["risk_counts"]
            if risk and sum(row["count"] for row in risk):
                axes[0, 1].pie([row["count"] for row in risk], labels=[row["label"] for row in risk], colors=["#B91C1C", "#D55E00", "#E69F00", "#0072B2"], startangle=90, wedgeprops={"width": .42, "edgecolor": "white"}, autopct=lambda p: f"{p:.0f}%" if p >= 4 else "")
            else:
                axes[0, 1].text(.5, .5, "No risk data", ha="center", va="center")
            axes[0, 1].set_title("Risk band composition")
            status = analytics["status_counts"]
            if status and sum(row["count"] for row in status):
                axes[1, 0].pie([row["count"] for row in status], labels=[row["label"] for row in status], colors=["#D55E00", "#0072B2", "#009E73", "#8B5CF6"], startangle=90, wedgeprops={"width": .42, "edgecolor": "white"}, autopct=lambda p: f"{p:.0f}%" if p >= 4 else "")
            else:
                axes[1, 0].text(.5, .5, "No status data", ha="center", va="center")
            axes[1, 0].set_title("Case status")
            top = analytics["top_parameters"][:8]
            axes[1, 1].barh([row["parameter"] for row in reversed(top)], [row["cases"] for row in reversed(top)], color="#009E73")
            axes[1, 1].set_title("Most frequent parameters")
            axes[1, 1].set_xlabel("Case count")
            axes[0, 0].set_ylabel("Case count")
            figure.suptitle(f"{data['water_type_label']} non-compliance · {data['start_date']} to {data['end_date']}")
            output = io.BytesIO()
            figure.savefig(output, format=export_type, dpi=180, bbox_inches="tight")
            plt.close(figure)
            return self.download(output.getvalue(), f"{filename}.{export_type}", "image/png" if export_type == "png" else "image/svg+xml")
        return HttpResponse("Unsupported export format.", status=400, content_type="text/plain")

    @staticmethod
    def download(content, filename, content_type):
        response = HttpResponse(content, content_type=content_type)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        return response
