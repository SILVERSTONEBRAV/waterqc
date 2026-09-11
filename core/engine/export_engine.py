import csv
import io
import json
import logging
from typing import Callable, Dict, Any, Iterable, Optional

from django.http import HttpResponse, StreamingHttpResponse
from django.utils.encoding import smart_str
from django.core.exceptions import ImproperlyConfigured

from openpyxl import Workbook
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Table
from reportlab.lib import colors

from core.registry import registry

logger = logging.getLogger(__name__)


class ExportEngine:
    def __init__(self):
        self._formatters: Dict[str, Callable] = {}

        self.register("csv", self._csv)
        self.register("json", self._json)
        self.register("xlsx", self._xlsx)
        self.register("docx", self._docx)
        self.register("pdf", self._pdf)

    def register(self, name: str, handler: Callable):
        self._formatters[name.lower()] = handler

    def get(self, name: str) -> Optional[Callable]:
        return self._formatters.get(name.lower())

    def run(self, ctx: Dict[str, Any]):
        export_type = (ctx.get("type") or "").lower().strip()

        if not export_type:
            raise ImproperlyConfigured("Export type required")

        handler = self.get(export_type)
        if not handler:
            raise ImproperlyConfigured(f"Unsupported export type: {export_type}")

        qs = self._qs(ctx)

        return handler(qs, ctx)

    def _qs(self, ctx: Dict[str, Any]):
        view = ctx.get("view")
        model = ctx["model"]
        request = ctx["request"]
        entry = ctx["entry"]

        if view and hasattr(view, "_qs"):
            qs = view._qs(ctx)
        else:
            qs = registry.resolve_queryset(model, request, entry.config)

        stream = ctx.get("stream", True)

        if stream:
            qs._prefetch_related_lookups = ()

        return qs

    def _fields(self, model):
        return [f.name for f in model._meta.fields if not f.primary_key]

    def _iter(self, qs, chunk_size=2000):
        return qs.iterator(chunk_size=chunk_size)

    def _csv(self, qs, ctx):
        model = ctx["model"]
        fields = self._fields(model)

        def gen():
            yield ",".join(fields) + "\n"

            for obj in self._iter(qs):
                yield ",".join(
                    smart_str(getattr(obj, f, "")) for f in fields
                ) + "\n"

        response = StreamingHttpResponse(gen(), content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{model._meta.model_name}.csv"'
        return response

    def _json(self, qs, ctx):
        def gen():
            yield '{"data":['

            first = True

            for obj in self._iter(qs):
                row = registry.resolve_row(obj)

                payload = json.dumps({
                    "id": row["id"],
                    "fields": {k: v["value"] for k, v in row["fields"].items()}
                })

                if not first:
                    yield ","
                first = False

                yield payload

            yield "]}"

        return StreamingHttpResponse(gen(), content_type="application/json")

    def _xlsx(self, qs, ctx):
        model = ctx["model"]
        fields = self._fields(model)

        wb = Workbook()
        ws = wb.active

        ws.append(fields)

        for obj in qs:
            ws.append([getattr(obj, f, "") for f in fields])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return HttpResponse(
            buf,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{model._meta.model_name}.xlsx"'
            },
        )

    def _docx(self, qs, ctx):
        model = ctx["model"]
        fields = self._fields(model)

        doc = Document()
        doc.add_heading(model.__name__, 0)

        table = doc.add_table(rows=1, cols=len(fields))

        for i, f in enumerate(fields):
            table.rows[0].cells[i].text = f

        for obj in qs:
            row = table.add_row().cells
            for i, f in enumerate(fields):
                row[i].text = str(getattr(obj, f, ""))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        return HttpResponse(
            buf,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f'attachment; filename="{model._meta.model_name}.docx"'
            },
        )

    def _pdf(self, qs, ctx):
        model = ctx["model"]
        fields = self._fields(model)

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf)

        data = [fields]

        for obj in self._iter(qs):
            data.append([str(getattr(obj, f, "")) for f in fields])

        table = Table(data)
        table.setStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black)
        ])

        doc.build([table])

        buf.seek(0)

        return HttpResponse(
            buf,
            content_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{model._meta.model_name}.pdf"'
            },
        )


export_engine = ExportEngine()
