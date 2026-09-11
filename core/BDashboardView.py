from __future__ import annotations

import csv
import io
import json
import logging
import math
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from statistics import mean, median, pstdev
from typing import Any

import pandas as pd

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.db import models
from django.db.models import (
    Avg,
    Case,
    Count,
    ExpressionWrapper,
    F,
    FloatField,
    IntegerField,
    Max,
    Min,
    Q,
    Sum,
    Value,
    When,
)
from django.db.models.functions import Coalesce, TruncDate, TruncMonth
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.generic import TemplateView

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from xhtml2pdf import pisa

try:
    from core.mixins import WanjalaMixin
except ImportError:
    from .mixins import WanjalaMixin

from .models import (
    BatchTestResult,
    Chemical,
    ChemicalAnalysisProcess,
    ChemicalAnalysisRequest,
    ChemicalBatch,
    ChemicalCoA,
)

logger = logging.getLogger(__name__)
User = get_user_model()


class BDashboardView(WanjalaMixin, TemplateView):
    template_name = "ca/dashboard/index.html"

    model = ChemicalAnalysisProcess

    permission_required_view = "ca.view_chemicalanalysisprocess"

    default_paginate_by = None

    export_file_name = "chemical_analysis_dashboard"

    dashboard_days_default = 30
    dashboard_days_max = 3650

    max_table_rows = 200
    max_chart_points = 365

    STATUS_ORDER = [
        "CREATED",
        "RECEIVED",
        "REQUESTED",
        "CHECKLIST",
        "ASSIGNED",
        "ANALYSIS",
        "RETEST",
        "QA",
        "CAPA",
        "APPROVED",
        "REJECTED",
        "UTILIZED",
        "RETURNED",
        "COMPLETED",
    ]

    TERMINAL_STATUSES = {
        "APPROVED",
        "REJECTED",
        "UTILIZED",
        "RETURNED",
        "COMPLETED",
    }

    ACTIVE_STATUSES = {
        "CREATED",
        "RECEIVED",
        "REQUESTED",
        "CHECKLIST",
        "ASSIGNED",
        "ANALYSIS",
        "RETEST",
        "QA",
        "CAPA",
    }

    QA_STATUSES = {
        "QA",
        "CAPA",
        "RETEST",
    }

    SUCCESS_BATCH_STATUSES = {
        "APPROVED",
        "UTILIZED",
    }

    FAILURE_BATCH_STATUSES = {
        "REJECTED",
        "RETURNED",
    }

    def dispatch(self, request, *args, **kwargs):
        self.request = request

        if request.GET.get("format") in {
            "json",
            "dashboard-json",
            "analytics",
        }:
            result = self.process_request(request)
            if result:
                return result

            self.check_idle_timeout()

            if not self.check_permission("view"):
                raise PermissionDenied

            return self.dashboard_json_response()

        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        qs = ChemicalAnalysisProcess.objects.all()

        qs = self._scope_queryset(qs)

        return qs

    def _scope_queryset(self, qs):
        user = self.request.user

        if self._is_superuser(user):
            return qs

        try:
            plants = self.get_user_plants()
        except Exception:
            plants = None

        if not plants:
            return qs.none()

        model = qs.model

        if self._has_field(model, "treatment_plant"):
            return qs.filter(treatment_plant__in=plants)

        if self._has_relation_path(model, "batch__treatment_plant"):
            return qs.filter(batch__treatment_plant__in=plants)

        if self._has_relation_path(model, "request__treatment_plant"):
            return qs.filter(request__treatment_plant__in=plants)

        return qs

    def _scope_related_queryset(self, qs):
        return self._scope_queryset(qs)

    @staticmethod
    def _has_field(model, field_name: str) -> bool:
        try:
            model._meta.get_field(field_name)
            return True
        except Exception:
            return False

    @classmethod
    def _has_relation_path(cls, model, path: str) -> bool:
        current = model

        for part in path.split("__"):
            try:
                field = current._meta.get_field(part)
            except Exception:
                return False

            remote = getattr(field, "remote_field", None)

            if remote is None or remote.model is None:
                return part == path.split("__")[-1]

            current = remote.model

        return True

    @staticmethod
    def _first_field(model, candidates):
        for candidate in candidates:
            try:
                model._meta.get_field(candidate)
                return candidate
            except Exception:
                continue
        return None

    @staticmethod
    def _field_names(model):
        return {
            field.name
            for field in model._meta.get_fields()
            if hasattr(field, "attname")
        }

    @staticmethod
    def _safe_get(obj, *paths, default=None):
        for path in paths:
            current = obj

            try:
                for part in path.split("__"):
                    current = getattr(current, part)
                    if callable(current):
                        current = current()

                if current is not None:
                    return current
            except Exception:
                continue

        return default

    @staticmethod
    def _string(value, default=""):
        if value is None:
            return default

        if isinstance(value, Decimal):
            return str(value)

        return str(value)

    @staticmethod
    def _number(value, default=0.0):
        try:
            if value is None:
                return default

            return float(value)
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _percent(numerator, denominator):
        if not denominator:
            return 0.0

        return round((numerator / denominator) * 100, 2)

    @staticmethod
    def _hours_between(start, end=None):
        if not start:
            return None

        end = end or timezone.now()

        try:
            return max(0.0, (end - start).total_seconds() / 3600)
        except Exception:
            return None

    def _date_range(self):
        today = timezone.localdate()

        period = self.request.GET.get("period", "30d")

        if period == "7d":
            days = 7
        elif period == "30d":
            days = 30
        elif period == "90d":
            days = 90
        elif period == "180d":
            days = 180
        elif period == "365d":
            days = 365
        elif period == "ytd":
            start = today.replace(month=1, day=1)
            return start, today
        elif period == "custom":
            try:
                start = datetime.strptime(
                    self.request.GET.get("start", ""),
                    "%Y-%m-%d",
                ).date()

                end = datetime.strptime(
                    self.request.GET.get("end", ""),
                    "%Y-%m-%d",
                ).date()

                if end < start:
                    start, end = end, start

                max_span = timedelta(days=self.dashboard_days_max)

                if end - start > max_span:
                    start = end - max_span

                return start, end
            except (TypeError, ValueError):
                days = self.dashboard_days_default
        else:
            days = self.dashboard_days_default

        days = min(days, self.dashboard_days_max)

        return today - timedelta(days=days - 1), today

    def _datetime_bounds(self):
        start, end = self._date_range()

        start_dt = timezone.make_aware(
            datetime.combine(start, datetime.min.time())
        )

        end_dt = timezone.make_aware(
            datetime.combine(
                end + timedelta(days=1),
                datetime.min.time(),
            )
        )

        return start_dt, end_dt

    def _date_field(self, model):
        return self._first_field(
            model,
            [
                "created_at",
                "received_at",
                "received_date",
                "created",
                "date_created",
                "timestamp",
                "analysis_date",
                "date",
            ],
        )

    def _status_field(self, model):
        return self._first_field(
            model,
            [
                "status",
                "state",
                "process_status",
            ],
        )

    def _user_field(self, model):
        return self._first_field(
            model,
            [
                "analyst",
                "assigned_analyst",
                "assigned_to",
                "performed_by",
                "user",
                "created_by",
            ],
        )

    def _chemical_path(self, model):
        if self._has_field(model, "chemical"):
            return "chemical"

        if self._has_relation_path(model, "batch__chemical"):
            return "batch__chemical"

        if self._has_relation_path(model, "request__chemical"):
            return "request__chemical"

        return None

    def _supplier_path(self, model):
        candidates = [
            "supplier",
            "batch__supplier",
            "request__supplier",
            "batch__chemical_supplier",
        ]

        for path in candidates:
            if self._has_relation_path(model, path):
                return path

        return None

    def _plant_path(self, model):
        candidates = [
            "treatment_plant",
            "batch__treatment_plant",
            "request__treatment_plant",
        ]

        for path in candidates:
            if self._has_relation_path(model, path):
                return path

        return None

    def _base_date_filtered_queryset(self, model):
        qs = model.objects.all()

        qs = self._scope_related_queryset(qs)

        date_field = self._date_field(model)

        if date_field:
            start_dt, end_dt = self._datetime_bounds()
            qs = qs.filter(
                **{
                    f"{date_field}__gte": start_dt,
                    f"{date_field}__lt": end_dt,
                }
            )

        return qs

    def _process_queryset(self):
        return self._base_date_filtered_queryset(
            ChemicalAnalysisProcess
        )

    def _batch_queryset(self):
        return self._base_date_filtered_queryset(
            ChemicalBatch
        )

    def _request_queryset(self):
        return self._base_date_filtered_queryset(
            ChemicalAnalysisRequest
        )

    def _result_queryset(self):
        qs = BatchTestResult.objects.all()

        qs = self._scope_related_queryset(qs)

        date_field = self._date_field(BatchTestResult)

        if date_field:
            start_dt, end_dt = self._datetime_bounds()

            qs = qs.filter(
                **{
                    f"{date_field}__gte": start_dt,
                    f"{date_field}__lt": end_dt,
                }
            )

        return qs

    def _coa_queryset(self):
        qs = ChemicalCoA.objects.all()

        qs = self._scope_related_queryset(qs)

        date_field = self._date_field(ChemicalCoA)

        if date_field:
            start_dt, end_dt = self._datetime_bounds()

            qs = qs.filter(
                **{
                    f"{date_field}__gte": start_dt,
                    f"{date_field}__lt": end_dt,
                }
            )

        return qs

    def _process_status_counts(self):
        qs = self._process_queryset()

        status_field = self._status_field(ChemicalAnalysisProcess)

        if not status_field:
            return {}

        rows = (
            qs.values(status_field)
            .annotate(total=Count("pk"))
            .order_by()
        )

        result = {
            status: 0
            for status in self.STATUS_ORDER
        }

        for row in rows:
            status = row.get(status_field)

            if status is not None:
                result[str(status).upper()] = row["total"]

        return result

    def _batch_status_counts(self):
        qs = self._batch_queryset()

        status_field = self._status_field(ChemicalBatch)

        if not status_field:
            return {}

        rows = (
            qs.values(status_field)
            .annotate(total=Count("pk"))
            .order_by()
        )

        return {
            str(row.get(status_field)).upper(): row["total"]
            for row in rows
            if row.get(status_field) is not None
        }

    def _request_status_counts(self):
        qs = self._request_queryset()

        status_field = self._status_field(
            ChemicalAnalysisRequest
        )

        if not status_field:
            return {}

        rows = (
            qs.values(status_field)
            .annotate(total=Count("pk"))
            .order_by()
        )

        return {
            str(row.get(status_field)).upper(): row["total"]
            for row in rows
            if row.get(status_field) is not None
        }

    def _result_status_counts(self):
        qs = self._result_queryset()

        status_field = self._status_field(BatchTestResult)

        if not status_field:
            return {}

        rows = (
            qs.values(status_field)
            .annotate(total=Count("pk"))
            .order_by()
        )

        return {
            str(row.get(status_field)).upper(): row["total"]
            for row in rows
            if row.get(status_field) is not None
        }

    def _basic_kpis(self):
        processes = self._process_queryset()
        batches = self._batch_queryset()
        requests = self._request_queryset()
        results = self._result_queryset()
        coas = self._coa_queryset()

        process_status = self._process_status_counts()
        batch_status = self._batch_status_counts()
        request_status = self._request_status_counts()
        result_status = self._result_status_counts()

        total_batches = batches.count()
        total_processes = processes.count()
        total_requests = requests.count()
        total_results = results.count()

        active_processes = sum(
            process_status.get(status, 0)
            for status in self.ACTIVE_STATUSES
        )

        qa_queue = sum(
            process_status.get(status, 0)
            for status in self.QA_STATUSES
        )

        passed = (
            result_status.get("PASS", 0)
            + batch_status.get("APPROVED", 0)
            + batch_status.get("UTILIZED", 0)
        )

        failed = (
            result_status.get("FAIL", 0)
            + batch_status.get("REJECTED", 0)
            + batch_status.get("RETURNED", 0)
        )

        completed = (
            process_status.get("COMPLETED", 0)
            + process_status.get("APPROVED", 0)
            + process_status.get("REJECTED", 0)
        )

        quality_denominator = passed + failed

        return {
            "total_batches": total_batches,
            "total_processes": total_processes,
            "total_requests": total_requests,
            "total_results": total_results,
            "active_analyses": active_processes,
            "pending_requests": request_status.get("PENDING", 0),
            "in_progress_requests": request_status.get("IN_PROGRESS", 0),
            "completed_requests": request_status.get("COMPLETED", 0),
            "qa_queue": qa_queue,
            "retest_queue": process_status.get("RETEST", 0),
            "capa_queue": process_status.get("CAPA", 0),
            "approved": batch_status.get("APPROVED", 0),
            "rejected": batch_status.get("REJECTED", 0),
            "utilized": batch_status.get("UTILIZED", 0),
            "returned": batch_status.get("RETURNED", 0),
            "completed": completed,
            "pass_count": passed,
            "fail_count": failed,
            "pass_rate": self._percent(
                passed,
                quality_denominator,
            ),
            "fail_rate": self._percent(
                failed,
                quality_denominator,
            ),
            "coa_count": coas.count(),
        }

    def _request_sla_metrics(self):
        qs = self._request_queryset()

        deadline_field = self._first_field(
            ChemicalAnalysisRequest,
            [
                "deadline",
                "due_at",
                "due_date",
                "sla_deadline",
            ],
        )

        completed_field = self._first_field(
            ChemicalAnalysisRequest,
            [
                "completed_at",
                "completed_on",
                "completion_date",
                "updated_at",
            ],
        )

        if not deadline_field:
            return {
                "deadline_field": None,
                "tracked": 0,
                "overdue": 0,
                "within_sla": 0,
                "sla_rate": 0,
                "overdue_rate": 0,
            }

        now = timezone.now()

        overdue = qs.filter(
            **{
                f"{deadline_field}__lt": now,
            }
        ).exclude(
            **{
                "status__in": [
                    "COMPLETED",
                    "APPROVED",
                    "REJECTED",
                ]
            }
        ).count()

        tracked = qs.exclude(
            **{
                f"{deadline_field}__isnull": True,
            }
        ).count()

        within_sla = max(tracked - overdue, 0)

        completed_on_time = 0

        if completed_field:
            completed_qs = qs.filter(
                **{
                    f"{completed_field}__isnull": False,
                    f"{deadline_field}__isnull": False,
                }
            )

            completed_on_time = completed_qs.filter(
                **{
                    f"{completed_field}__lte": F(deadline_field),
                }
            ).count()

        effective_within_sla = (
            completed_on_time
            if completed_field
            else within_sla
        )

        effective_tracked = (
            completed_qs.count()
            if completed_field
            else tracked
        ) if completed_field else tracked

        return {
            "deadline_field": deadline_field,
            "tracked": tracked,
            "overdue": overdue,
            "within_sla": within_sla,
            "completed_on_time": completed_on_time,
            "sla_rate": self._percent(
                effective_within_sla,
                effective_tracked,
            ),
            "overdue_rate": self._percent(
                overdue,
                tracked,
            ),
        }

    def _turnaround_metrics(self):
        process_model = ChemicalAnalysisProcess

        start_field = self._first_field(
            process_model,
            [
                "analysis_started_at",
                "started_at",
                "analysis_start",
                "received_at",
                "created_at",
            ],
        )

        end_field = self._first_field(
            process_model,
            [
                "analysis_completed_at",
                "completed_at",
                "analysis_end",
                "updated_at",
            ],
        )

        if not start_field or not end_field:
            return {
                "available": False,
                "average_hours": 0,
                "median_hours": 0,
                "p90_hours": 0,
                "min_hours": 0,
                "max_hours": 0,
            }

        qs = self._process_queryset().filter(
            **{
                f"{start_field}__isnull": False,
                f"{end_field}__isnull": False,
            }
        ).values(
            start_field,
            end_field,
        )[:5000]

        durations = []

        for row in qs:
            start = row.get(start_field)
            end = row.get(end_field)

            hours = self._hours_between(start, end)

            if hours is not None and hours < 10000:
                durations.append(hours)

        if not durations:
            return {
                "available": True,
                "average_hours": 0,
                "median_hours": 0,
                "p90_hours": 0,
                "min_hours": 0,
                "max_hours": 0,
            }

        sorted_values = sorted(durations)

        p90_index = min(
            len(sorted_values) - 1,
            max(0, math.ceil(len(sorted_values) * 0.90) - 1),
        )

        return {
            "available": True,
            "sample_size": len(durations),
            "average_hours": round(mean(durations), 2),
            "median_hours": round(median(durations), 2),
            "p90_hours": round(sorted_values[p90_index], 2),
            "min_hours": round(min(durations), 2),
            "max_hours": round(max(durations), 2),
        }

    def _daily_process_trend(self):
        qs = self._process_queryset()

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        if not date_field:
            return []

        rows = (
            qs.annotate(
                day=TruncDate(date_field)
            )
            .values("day")
            .annotate(
                total=Count("pk"),
            )
            .order_by("day")
        )

        return [
            {
                "date": row["day"].isoformat()
                if row["day"]
                else None,
                "total": row["total"],
            }
            for row in rows
        ][-self.max_chart_points :]

    def _daily_quality_trend(self):
        qs = self._result_queryset()

        date_field = self._date_field(
            BatchTestResult
        )

        status_field = self._status_field(
            BatchTestResult
        )

        if not date_field or not status_field:
            return []

        rows = (
            qs.annotate(
                day=TruncDate(date_field)
            )
            .values("day")
            .annotate(
                total=Count("pk"),
                passed=Count(
                    "pk",
                    filter=Q(**{status_field: "PASS"}),
                ),
                failed=Count(
                    "pk",
                    filter=Q(**{status_field: "FAIL"}),
                ),
            )
            .order_by("day")
        )

        result = []

        for row in rows:
            denominator = row["passed"] + row["failed"]

            result.append(
                {
                    "date": (
                        row["day"].isoformat()
                        if row["day"]
                        else None
                    ),
                    "total": row["total"],
                    "passed": row["passed"],
                    "failed": row["failed"],
                    "pass_rate": self._percent(
                        row["passed"],
                        denominator,
                    ),
                }
            )

        return result[-self.max_chart_points :]

    def _monthly_process_trend(self):
        qs = self._process_queryset()

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        if not date_field:
            return []

        rows = (
            qs.annotate(
                month=TruncMonth(date_field)
            )
            .values("month")
            .annotate(
                total=Count("pk"),
            )
            .order_by("month")
        )

        return [
            {
                "month": row["month"].strftime("%Y-%m")
                if row["month"]
                else None,
                "total": row["total"],
            }
            for row in rows
        ]

    def _workflow_funnel(self):
        counts = self._process_status_counts()

        stages = [
            ("Received", ["RECEIVED"]),
            ("Requested", ["REQUESTED"]),
            ("Assigned", ["ASSIGNED"]),
            ("Checklist", ["CHECKLIST"]),
            ("Analysis", ["ANALYSIS"]),
            ("QA", ["QA", "RETEST", "CAPA"]),
            ("Approved", ["APPROVED"]),
            ("Completed", ["COMPLETED", "UTILIZED"]),
        ]

        result = []

        previous = None

        for label, statuses in stages:
            value = sum(
                counts.get(status, 0)
                for status in statuses
            )

            conversion = (
                self._percent(value, previous)
                if previous
                else 100.0
            )

            result.append(
                {
                    "stage": label,
                    "count": value,
                    "conversion": conversion,
                }
            )

            if value:
                previous = value

        return result

    def _chemical_performance(self):
        model = ChemicalAnalysisProcess

        chemical_path = self._chemical_path(model)

        if not chemical_path:
            return []

        status_field = self._status_field(model)

        if not status_field:
            return []

        rows = (
            self._process_queryset()
            .values(
                f"{chemical_path}__id",
                f"{chemical_path}__code",
                f"{chemical_path}__name",
                status_field,
            )
            .annotate(
                total=Count("pk"),
            )
            .order_by()
        )

        grouped = defaultdict(
            lambda: {
                "total": 0,
                "approved": 0,
                "rejected": 0,
                "completed": 0,
            }
        )

        for row in rows:
            key = (
                row.get(f"{chemical_path}__id")
                or row.get(f"{chemical_path}__code")
                or row.get(f"{chemical_path}__name")
                or "Unknown"
            )

            status = str(
                row.get(status_field) or ""
            ).upper()

            grouped[str(key)]["total"] += row["total"]

            if status == "APPROVED":
                grouped[str(key)]["approved"] += row["total"]

            if status == "REJECTED":
                grouped[str(key)]["rejected"] += row["total"]

            if status == "COMPLETED":
                grouped[str(key)]["completed"] += row["total"]

        result = []

        for key, values in grouped.items():
            denominator = (
                values["approved"]
                + values["rejected"]
            )

            result.append(
                {
                    "chemical": key,
                    **values,
                    "quality_rate": self._percent(
                        values["approved"],
                        denominator,
                    ),
                    "risk_score": self._risk_score(
                        failure_count=values["rejected"],
                        active_count=(
                            values["total"]
                            - values["completed"]
                        ),
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: (
                -x["risk_score"],
                -x["total"],
            ),
        )[:50]

    def _supplier_performance(self):
        model = ChemicalAnalysisProcess

        supplier_path = self._supplier_path(model)

        if not supplier_path:
            return []

        status_field = self._status_field(model)

        if not status_field:
            return []

        rows = (
            self._process_queryset()
            .values(
                f"{supplier_path}__id",
                f"{supplier_path}__name",
                f"{supplier_path}__company_name",
                status_field,
            )
            .annotate(
                total=Count("pk"),
            )
            .order_by()
        )

        grouped = defaultdict(
            lambda: {
                "total": 0,
                "approved": 0,
                "rejected": 0,
            }
        )

        names = {}

        for row in rows:
            supplier_id = (
                row.get(f"{supplier_path}__id")
                or "unknown"
            )

            supplier_name = (
                row.get(f"{supplier_path}__name")
                or row.get(
                    f"{supplier_path}__company_name"
                )
                or str(supplier_id)
            )

            names[str(supplier_id)] = supplier_name

            status = str(
                row.get(status_field) or ""
            ).upper()

            grouped[str(supplier_id)]["total"] += row["total"]

            if status in {"APPROVED", "UTILIZED"}:
                grouped[str(supplier_id)]["approved"] += row["total"]

            elif status in {"REJECTED", "RETURNED"}:
                grouped[str(supplier_id)]["rejected"] += row["total"]

        result = []

        for supplier_id, values in grouped.items():
            denominator = (
                values["approved"]
                + values["rejected"]
            )

            result.append(
                {
                    "supplier": names.get(
                        supplier_id,
                        supplier_id,
                    ),
                    **values,
                    "acceptance_rate": self._percent(
                        values["approved"],
                        denominator,
                    ),
                    "risk_score": self._risk_score(
                        failure_count=values["rejected"],
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: (
                -x["risk_score"],
                -x["total"],
            ),
        )[:50]

    def _analyst_performance(self):
        model = ChemicalAnalysisProcess

        user_field = self._user_field(model)

        if not user_field:
            return []

        status_field = self._status_field(model)

        if not status_field:
            return []

        rows = (
            self._process_queryset()
            .values(
                f"{user_field}__id",
                f"{user_field}__username",
                f"{user_field}__first_name",
                f"{user_field}__last_name",
                status_field,
            )
            .annotate(
                total=Count("pk"),
            )
            .order_by()
        )

        grouped = defaultdict(
            lambda: {
                "total": 0,
                "completed": 0,
                "retest": 0,
                "capa": 0,
            }
        )

        names = {}

        for row in rows:
            user_id = (
                row.get(f"{user_field}__id")
                or "unknown"
            )

            username = row.get(
                f"{user_field}__username"
            )

            full_name = " ".join(
                filter(
                    None,
                    [
                        row.get(
                            f"{user_field}__first_name"
                        ),
                        row.get(
                            f"{user_field}__last_name"
                        ),
                    ],
                )
            )

            names[str(user_id)] = (
                full_name
                or username
                or str(user_id)
            )

            status = str(
                row.get(status_field) or ""
            ).upper()

            grouped[str(user_id)]["total"] += row["total"]

            if status in {
                "COMPLETED",
                "APPROVED",
                "UTILIZED",
            }:
                grouped[str(user_id)]["completed"] += row["total"]

            if status == "RETEST":
                grouped[str(user_id)]["retest"] += row["total"]

            if status == "CAPA":
                grouped[str(user_id)]["capa"] += row["total"]

        result = []

        for user_id, values in grouped.items():
            result.append(
                {
                    "analyst": names.get(
                        user_id,
                        user_id,
                    ),
                    **values,
                    "completion_rate": self._percent(
                        values["completed"],
                        values["total"],
                    ),
                    "quality_risk": self._risk_score(
                        failure_count=(
                            values["retest"]
                            + values["capa"]
                        ),
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: (
                -x["completed"],
                x["quality_risk"],
            ),
        )[:50]

    def _failure_pareto(self):
        model = ChemicalAnalysisProcess

        chemical_path = self._chemical_path(model)

        if not chemical_path:
            return []

        status_field = self._status_field(model)

        if not status_field:
            return []

        rows = (
            self._process_queryset()
            .filter(
                **{
                    f"{status_field}__in": [
                        "REJECTED",
                        "RETURNED",
                    ]
                }
            )
            .values(
                f"{chemical_path}__id",
                f"{chemical_path}__code",
                f"{chemical_path}__name",
            )
            .annotate(
                failures=Count("pk")
            )
            .order_by("-failures")
        )

        values = [
            {
                "chemical": (
                    row.get(
                        f"{chemical_path}__code"
                    )
                    or row.get(
                        f"{chemical_path}__name"
                    )
                    or str(
                        row.get(
                            f"{chemical_path}__id"
                        )
                    )
                ),
                "failures": row["failures"],
            }
            for row in rows
        ]

        total = sum(
            row["failures"]
            for row in values
        )

        cumulative = 0

        for row in values:
            cumulative += row["failures"]

            row["cumulative_percent"] = self._percent(
                cumulative,
                total,
            )

        return values[:50]

    def _aging_analysis(self):
        qs = self._process_queryset()

        status_field = self._status_field(
            ChemicalAnalysisProcess
        )

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        if not status_field or not date_field:
            return []

        active_qs = qs.exclude(
            **{
                f"{status_field}__in": list(
                    self.TERMINAL_STATUSES
                )
            }
        ).values(
            "pk",
            status_field,
            date_field,
        ).order_by(
            date_field
        )[:self.max_table_rows]

        result = []

        now = timezone.now()

        for row in active_qs:
            created = row.get(date_field)

            if not created:
                continue

            hours = self._hours_between(
                created,
                now,
            )

            if hours is None:
                continue

            risk = self._aging_risk(hours)

            result.append(
                {
                    "id": row["pk"],
                    "status": row.get(status_field),
                    "age_hours": round(hours, 2),
                    "age_days": round(
                        hours / 24,
                        2,
                    ),
                    "risk": risk,
                }
            )

        return sorted(
            result,
            key=lambda x: -x["age_hours"],
        )

    @staticmethod
    def _aging_risk(hours):
        if hours >= 72:
            return "CRITICAL"

        if hours >= 48:
            return "HIGH"

        if hours >= 24:
            return "MEDIUM"

        return "LOW"

    @staticmethod
    def _risk_score(
        failure_count=0,
        active_count=0,
        overdue_count=0,
        rework_count=0,
    ):
        score = (
            failure_count * 10
            + active_count * 2
            + overdue_count * 15
            + rework_count * 8
        )

        return min(
            100,
            round(score),
        )

    def _process_risk_rows(self):
        qs = self._process_queryset()

        status_field = self._status_field(
            ChemicalAnalysisProcess
        )

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        if not status_field:
            return []

        fields = ["pk", status_field]

        if date_field:
            fields.append(date_field)

        rework_field = self._first_field(
            ChemicalAnalysisProcess,
            [
                "rework_count",
                "rework",
                "retry_count",
            ],
        )

        escalation_field = self._first_field(
            ChemicalAnalysisProcess,
            [
                "escalation_count",
                "escalations",
            ],
        )

        if rework_field:
            fields.append(rework_field)

        if escalation_field:
            fields.append(escalation_field)

        rows = qs.values(*fields).order_by(
            f"-{date_field}"
            if date_field
            else "-pk"
        )[:self.max_table_rows]

        result = []

        now = timezone.now()

        for row in rows:
            status = str(
                row.get(status_field) or ""
            ).upper()

            age_hours = (
                self._hours_between(
                    row.get(date_field),
                    now,
                )
                if date_field
                else 0
            ) or 0

            rework = (
                self._number(
                    row.get(rework_field)
                )
                if rework_field
                else 0
            )

            escalation = (
                self._number(
                    row.get(escalation_field)
                )
                if escalation_field
                else 0
            )

            score = self._risk_score(
                active_count=(
                    1
                    if status in self.ACTIVE_STATUSES
                    else 0
                ),
                rework_count=rework,
                overdue_count=(
                    1
                    if age_hours >= 24
                    and status in self.ACTIVE_STATUSES
                    else 0
                ),
            )

            score += int(escalation * 10)

            result.append(
                {
                    "id": row["pk"],
                    "status": status,
                    "age_hours": round(
                        age_hours,
                        2,
                    ),
                    "rework_count": int(rework),
                    "escalation_count": int(
                        escalation
                    ),
                    "risk_score": min(
                        score,
                        100,
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: -x["risk_score"],
        )

    def _detect_anomalies(self, daily_trend):
        if len(daily_trend) < 7:
            return []

        values = [
            self._number(
                row.get("total")
            )
            for row in daily_trend
        ]

        baseline = values[:-1]

        if not baseline:
            return []

        avg = mean(baseline)

        deviation = pstdev(
            baseline
        ) if len(baseline) > 1 else 0

        if deviation == 0:
            threshold = avg * 2
        else:
            threshold = avg + (2 * deviation)

        anomalies = []

        for row, value in zip(
            daily_trend,
            values,
        ):
            if value > threshold:
                anomalies.append(
                    {
                        "date": row["date"],
                        "value": value,
                        "baseline": round(
                            avg,
                            2,
                        ),
                        "type": "VOLUME_SPIKE",
                        "severity": "HIGH",
                    }
                )

        return anomalies[-30:]

    def _forecast(self, daily_trend):
        if len(daily_trend) < 7:
            return {
                "available": False,
                "horizon": 0,
                "values": [],
            }

        values = [
            self._number(
                row.get("total")
            )
            for row in daily_trend
        ]

        n = len(values)

        x_mean = (
            (n - 1) / 2
        )

        y_mean = mean(values)

        numerator = sum(
            (i - x_mean) * (y - y_mean)
            for i, y in enumerate(values)
        )

        denominator = sum(
            (i - x_mean) ** 2
            for i in range(n)
        )

        slope = (
            numerator / denominator
            if denominator
            else 0
        )

        intercept = (
            y_mean
            - slope * x_mean
        )

        last_date = (
            datetime.fromisoformat(
                daily_trend[-1]["date"]
            ).date()
        )

        forecast = []

        horizon = 7

        for i in range(1, horizon + 1):
            index = n - 1 + i

            predicted = max(
                0,
                intercept + slope * index,
            )

            forecast.append(
                {
                    "date": (
                        last_date
                        + timedelta(days=i)
                    ).isoformat(),
                    "predicted": round(
                        predicted,
                        2,
                    ),
                }
            )

        return {
            "available": True,
            "method": "linear_trend",
            "horizon": horizon,
            "values": forecast,
        }

    def _active_process_table(self):
        qs = self._process_queryset()

        status_field = self._status_field(
            ChemicalAnalysisProcess
        )

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        chemical_path = self._chemical_path(
            ChemicalAnalysisProcess
        )

        fields = ["pk"]

        if status_field:
            fields.append(status_field)

        if date_field:
            fields.append(date_field)

        if chemical_path:
            fields.extend(
                [
                    f"{chemical_path}__code",
                    f"{chemical_path}__name",
                ]
            )

        rows = qs.exclude(
            **{
                f"{status_field}__in": list(
                    self.TERMINAL_STATUSES
                )
            }
        ).values(
            *fields
        )

        result = []

        now = timezone.now()

        for row in rows[:self.max_table_rows]:
            created = (
                row.get(date_field)
                if date_field
                else None
            )

            age_hours = (
                self._hours_between(
                    created,
                    now,
                )
                if created
                else 0
            ) or 0

            chemical = (
                row.get(
                    f"{chemical_path}__code"
                )
                or row.get(
                    f"{chemical_path}__name"
                )
                if chemical_path
                else ""
            )

            result.append(
                {
                    "id": row["pk"],
                    "chemical": chemical or "—",
                    "status": (
                        row.get(status_field)
                        if status_field
                        else ""
                    ),
                    "created_at": (
                        created.isoformat()
                        if created
                        else None
                    ),
                    "age_hours": round(
                        age_hours,
                        2,
                    ),
                    "risk": self._aging_risk(
                        age_hours
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: -x["age_hours"],
        )

    def _qa_queue(self):
        qs = self._process_queryset()

        status_field = self._status_field(
            ChemicalAnalysisProcess
        )

        if not status_field:
            return []

        qs = qs.filter(
            **{
                f"{status_field}__in": list(
                    self.QA_STATUSES
                )
            }
        )

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        fields = ["pk", status_field]

        if date_field:
            fields.append(date_field)

        rows = qs.values(
            *fields
        ).order_by(
            f"-{date_field}"
            if date_field
            else "-pk"
        )[:self.max_table_rows]

        now = timezone.now()

        return [
            {
                "id": row["pk"],
                "status": row.get(
                    status_field
                ),
                "age_hours": round(
                    self._hours_between(
                        row.get(date_field),
                        now,
                    ) or 0,
                    2,
                ),
            }
            for row in rows
        ]

    def _approval_queue(self):
        qs = self._process_queryset()

        status_field = self._status_field(
            ChemicalAnalysisProcess
        )

        if not status_field:
            return []

        rows = qs.filter(
            **{
                f"{status_field}__in": [
                    "APPROVED",
                    "QA",
                ]
            }
        )

        date_field = self._date_field(
            ChemicalAnalysisProcess
        )

        fields = ["pk", status_field]

        if date_field:
            fields.append(date_field)

        rows = rows.values(
            *fields
        ).order_by(
            f"-{date_field}"
            if date_field
            else "-pk"
        )[:self.max_table_rows]

        return [
            {
                "id": row["pk"],
                "status": row.get(
                    status_field
                ),
                "date": (
                    row.get(date_field).isoformat()
                    if date_field
                    and row.get(date_field)
                    else None
                ),
            }
            for row in rows
        ]

    def _plant_summary(self):
        if self._is_superuser(
            self.request.user
        ):
            plants = self.get_user_plants()
        else:
            plants = self.get_user_plants()

        result = []

        process_model = ChemicalAnalysisProcess

        plant_path = self._plant_path(
            process_model
        )

        status_field = self._status_field(
            process_model
        )

        if not plant_path or not status_field:
            return result

        rows = (
            self._process_queryset()
            .values(
                f"{plant_path}__id",
                f"{plant_path}__name",
                status_field,
            )
            .annotate(
                total=Count("pk")
            )
            .order_by()
        )

        grouped = defaultdict(
            lambda: {
                "total": 0,
                "active": 0,
                "completed": 0,
                "rejected": 0,
            }
        )

        names = {}

        for row in rows:
            plant_id = row.get(
                f"{plant_path}__id"
            )

            key = str(
                plant_id or "unknown"
            )

            names[key] = (
                row.get(
                    f"{plant_path}__name"
                )
                or key
            )

            total = row["total"]

            status = str(
                row.get(status_field) or ""
            ).upper()

            grouped[key]["total"] += total

            if status in self.ACTIVE_STATUSES:
                grouped[key]["active"] += total

            if status in {
                "COMPLETED",
                "APPROVED",
                "UTILIZED",
            }:
                grouped[key]["completed"] += total

            if status in {
                "REJECTED",
                "RETURNED",
            }:
                grouped[key]["rejected"] += total

        allowed_ids = {
            str(plant.pk)
            for plant in plants
        }

        for key, values in grouped.items():
            if (
                not self._is_superuser(
                    self.request.user
                )
                and key not in allowed_ids
            ):
                continue

            result.append(
                {
                    "plant": names.get(
                        key,
                        key,
                    ),
                    **values,
                    "completion_rate": self._percent(
                        values["completed"],
                        values["total"],
                    ),
                    "rejection_rate": self._percent(
                        values["rejected"],
                        values["total"],
                    ),
                }
            )

        return sorted(
            result,
            key=lambda x: -x["total"],
        )

    def _build_charts(self):
        daily = self._daily_process_trend()
        quality = self._daily_quality_trend()
        workflow = self._workflow_funnel()
        chemical = self._chemical_performance()
        supplier = self._supplier_performance()
        pareto = self._failure_pareto()
        monthly = self._monthly_process_trend()

        return {
            "volume": {
                "type": "line",
                "labels": [
                    row["date"]
                    for row in daily
                ],
                "datasets": [
                    {
                        "label": "Analyses",
                        "data": [
                            row["total"]
                            for row in daily
                        ],
                    }
                ],
            },
            "quality": {
                "type": "line",
                "labels": [
                    row["date"]
                    for row in quality
                ],
                "datasets": [
                    {
                        "label": "Pass rate %",
                        "data": [
                            row["pass_rate"]
                            for row in quality
                        ],
                    },
                    {
                        "label": "Passed",
                        "data": [
                            row["passed"]
                            for row in quality
                        ],
                    },
                    {
                        "label": "Failed",
                        "data": [
                            row["failed"]
                            for row in quality
                        ],
                    },
                ],
            },
            "workflow": {
                "type": "bar",
                "labels": [
                    row["stage"]
                    for row in workflow
                ],
                "datasets": [
                    {
                        "label": "Workflow volume",
                        "data": [
                            row["count"]
                            for row in workflow
                        ],
                    }
                ],
            },
            "monthly": {
                "type": "bar",
                "labels": [
                    row["month"]
                    for row in monthly
                ],
                "datasets": [
                    {
                        "label": "Analyses",
                        "data": [
                            row["total"]
                            for row in monthly
                        ],
                    }
                ],
            },
            "chemical": {
                "type": "bar",
                "labels": [
                    row["chemical"]
                    for row in chemical[:20]
                ],
                "datasets": [
                    {
                        "label": "Quality rate %",
                        "data": [
                            row["quality_rate"]
                            for row in chemical[:20]
                        ],
                    }
                ],
            },
            "suppliers": {
                "type": "bar",
                "labels": [
                    row["supplier"]
                    for row in supplier[:20]
                ],
                "datasets": [
                    {
                        "label": "Acceptance rate %",
                        "data": [
                            row["acceptance_rate"]
                            for row in supplier[:20]
                        ],
                    }
                ],
            },
            "failures": {
                "type": "bar",
                "labels": [
                    row["chemical"]
                    for row in pareto[:20]
                ],
                "datasets": [
                    {
                        "label": "Failures",
                        "data": [
                            row["failures"]
                            for row in pareto[:20]
                        ],
                    },
                    {
                        "label": "Cumulative %",
                        "data": [
                            row["cumulative_percent"]
                            for row in pareto[:20]
                        ],
                    },
                ],
            },
        }

    def get_dashboard_context(self):
        start, end = self._date_range()

        kpis = self._basic_kpis()
        sla = self._request_sla_metrics()
        turnaround = self._turnaround_metrics()

        daily = self._daily_process_trend()

        context = {
            "dashboard": {
                "title": "Chemical Analysis Command Center",
                "subtitle": "Quality, workflow, SLA and operational intelligence",
                "generated_at": timezone.now(),
            },
            "filters": {
                "period": self.request.GET.get(
                    "period",
                    "30d",
                ),
                "start": start.isoformat(),
                "end": end.isoformat(),
                "available_periods": [
                    "7d",
                    "30d",
                    "90d",
                    "180d",
                    "365d",
                    "ytd",
                    "custom",
                ],
            },
            "kpis": kpis,
            "sla": sla,
            "turnaround": turnaround,
            "workflow": self._workflow_funnel(),
            "process_status": self._process_status_counts(),
            "batch_status": self._batch_status_counts(),
            "request_status": self._request_status_counts(),
            "result_status": self._result_status_counts(),
            "chemical_performance": self._chemical_performance(),
            "supplier_performance": self._supplier_performance(),
            "analyst_performance": self._analyst_performance(),
            "plant_summary": self._plant_summary(),
            "failure_pareto": self._failure_pareto(),
            "aging": self._aging_analysis(),
            "risk_register": self._process_risk_rows(),
            "active_processes": self._active_process_table(),
            "qa_queue": self._qa_queue(),
            "approval_queue": self._approval_queue(),
            "daily_trend": daily,
            "quality_trend": self._daily_quality_trend(),
            "monthly_trend": self._monthly_process_trend(),
            "anomalies": self._detect_anomalies(daily),
            "forecast": self._forecast(daily),
            "charts": self._build_charts(),
            "permissions": self._dashboard_permissions(),
            "urls": self._dashboard_urls(),
            "refresh": {
                "interval_seconds": 60,
                "server_time": timezone.now().isoformat(),
            },
        }

        return context

    def _dashboard_permissions(self):
        user = self.request.user

        permissions = {
            "view": user.has_perm(
                "ca.view_chemicalanalysisprocess"
            ),
            "add_batch": user.has_perm(
                "ca.add_chemicalbatch"
            ),
            "receive_batch": user.has_perm(
                "ca.receive_chemical_batch"
            ),
            "create_request": user.has_perm(
                "ca.add_chemicalanalysisrequest"
            ),
            "assign_analyst": user.has_perm(
                "ca.assign_chemical_analyst"
            ),
            "generate_coa": user.has_perm(
                "ca.generate_chemical_coa"
            ),
            "approve_coa": user.has_perm(
                "ca.approve_chemical_coa"
            ),
            "utilize_batch": user.has_perm(
                "ca.utilize_chemical_batch"
            ),
            "export": user.has_perm(
                "ca.export_chemicalanalysisprocess"
            ),
        }

        return permissions

    def _reverse_safe(self, name, fallback="#"):
        from django.urls import NoReverseMatch, reverse

        try:
            return reverse(name)
        except NoReverseMatch:
            return fallback

    def _dashboard_urls(self):
        return {
            "dashboard": self._reverse_safe(
                "ca:dashboard"
            ),
            "batches": self._reverse_safe(
                "ca:chemical_batch_list"
            ),
            "requests": self._reverse_safe(
                "ca:chemical_analysis_request_list"
            ),
            "processes": self._reverse_safe(
                "ca:process_list"
            ),
            "analytics": (
                f"{self.request.path}"
                "?format=json"
            ),
            "export_csv": (
                f"{self.request.path}"
                "?export=csv"
            ),
            "export_xlsx": (
                f"{self.request.path}"
                "?export=xlsx"
            ),
            "export_pdf": (
                f"{self.request.path}"
                "?export=pdf"
            ),
            "export_docx": (
                f"{self.request.path}"
                "?export=docx"
            ),
        }

    def get_context_data(self, **kwargs):
        context = super().get_context_data(
            **kwargs
        )

        context.update(
            self.get_dashboard_context()
        )

        return context

    def render_to_response(
        self,
        context,
        **response_kwargs,
    ):
        if self.is_partial():
            template = (
                "ca/dashboard/partials/dashboard.html"
            )

            return render(
                self.request,
                template,
                context,
                **response_kwargs,
            )

        return super().render_to_response(
            context,
            **response_kwargs,
        )

    def dashboard_json_response(self):
        context = self.get_dashboard_context()

        return JsonResponse(
            context,
            safe=False,
        )

    def _export_rows(self):
        context = self.get_dashboard_context()

        rows = []

        for row in context.get(
            "active_processes",
            [],
        ):
            rows.append(
                {
                    "process_id": row.get("id"),
                    "chemical": row.get(
                        "chemical"
                    ),
                    "status": row.get(
                        "status"
                    ),
                    "age_hours": row.get(
                        "age_hours"
                    ),
                    "risk": row.get(
                        "risk"
                    ),
                }
            )

        return context, rows

    def export_dashboard_csv(self):
        context, rows = self._export_rows()

        response = HttpResponse(
            content_type="text/csv; charset=utf-8"
        )

        filename = self.get_export_file_name(
            "csv"
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; filename="{filename}"'
        )

        writer = csv.writer(response)

        writer.writerow(
            [
                "Chemical Analysis Dashboard"
            ]
        )

        writer.writerow(
            [
                "Period",
                context["filters"]["start"],
                context["filters"]["end"],
            ]
        )

        writer.writerow([])

        writer.writerow(
            [
                "KPI",
                "Value",
            ]
        )

        for key, value in context[
            "kpis"
        ].items():
            writer.writerow(
                [
                    key,
                    value,
                ]
            )

        writer.writerow([])

        writer.writerow(
            [
                "Process ID",
                "Chemical",
                "Status",
                "Age Hours",
                "Risk",
            ]
        )

        for row in rows:
            writer.writerow(
                [
                    row["process_id"],
                    row["chemical"],
                    row["status"],
                    row["age_hours"],
                    row["risk"],
                ]
            )

        return response

    def export_dashboard_excel(self):
        context = self.get_dashboard_context()

        wb = Workbook()

        summary = wb.active
        summary.title = "Dashboard"

        summary.append(
            [
                "Chemical Analysis Command Center"
            ]
        )

        summary.append(
            [
                "Period",
                context["filters"]["start"],
                context["filters"]["end"],
            ]
        )

        summary.append([])

        summary.append(
            [
                "KPI",
                "Value",
            ]
        )

        for key, value in context[
            "kpis"
        ].items():
            summary.append(
                [
                    key,
                    value,
                ]
            )

        self._style_excel_sheet(
            summary
        )

        active = wb.create_sheet(
            "Active Processes"
        )

        active_headers = [
            "Process ID",
            "Chemical",
            "Status",
            "Age Hours",
            "Risk",
        ]

        active.append(active_headers)

        for row in context[
            "active_processes"
        ]:
            active.append(
                [
                    row.get("id"),
                    row.get("chemical"),
                    row.get("status"),
                    row.get("age_hours"),
                    row.get("risk"),
                ]
            )

        self._style_excel_sheet(active)

        chemicals = wb.create_sheet(
            "Chemical Performance"
        )

        chemical_headers = [
            "Chemical",
            "Total",
            "Approved",
            "Rejected",
            "Quality Rate",
            "Risk Score",
        ]

        chemicals.append(
            chemical_headers
        )

        for row in context[
            "chemical_performance"
        ]:
            chemicals.append(
                [
                    row.get("chemical"),
                    row.get("total"),
                    row.get("approved"),
                    row.get("rejected"),
                    row.get("quality_rate"),
                    row.get("risk_score"),
                ]
            )

        self._style_excel_sheet(
            chemicals
        )

        suppliers = wb.create_sheet(
            "Supplier Performance"
        )

        supplier_headers = [
            "Supplier",
            "Total",
            "Approved",
            "Rejected",
            "Acceptance Rate",
            "Risk Score",
        ]

        suppliers.append(
            supplier_headers
        )

        for row in context[
            "supplier_performance"
        ]:
            suppliers.append(
                [
                    row.get("supplier"),
                    row.get("total"),
                    row.get("approved"),
                    row.get("rejected"),
                    row.get(
                        "acceptance_rate"
                    ),
                    row.get("risk_score"),
                ]
            )

        self._style_excel_sheet(
            suppliers
        )

        analysts = wb.create_sheet(
            "Analyst Performance"
        )

        analyst_headers = [
            "Analyst",
            "Total",
            "Completed",
            "Retest",
            "CAPA",
            "Completion Rate",
            "Risk",
        ]

        analysts.append(
            analyst_headers
        )

        for row in context[
            "analyst_performance"
        ]:
            analysts.append(
                [
                    row.get("analyst"),
                    row.get("total"),
                    row.get("completed"),
                    row.get("retest"),
                    row.get("capa"),
                    row.get(
                        "completion_rate"
                    ),
                    row.get(
                        "quality_risk"
                    ),
                ]
            )

        self._style_excel_sheet(
            analysts
        )

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )

        filename = self.get_export_file_name(
            "xlsx"
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; filename="{filename}"'
        )

        wb.save(response)

        return response

    @staticmethod
    def _style_excel_sheet(ws):
        for cell in ws[1]:
            cell.font = Font(
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for column in ws.columns:
            max_length = 0

            column_letter = column[0].column_letter

            for cell in column:
                try:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )
                except Exception:
                    pass

            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                45,
            )

    def export_dashboard_docx(self):
        context = self.get_dashboard_context()

        document = Document()

        document.add_heading(
            "Chemical Analysis Command Center",
            level=1,
        )

        document.add_paragraph(
            f"Reporting period: "
            f"{context['filters']['start']} "
            f"to "
            f"{context['filters']['end']}"
        )

        document.add_heading(
            "Executive KPIs",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"

        for key, value in context[
            "kpis"
        ].items():
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value)

        document.add_heading(
            "Chemical Performance",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=6,
        )

        headers = [
            "Chemical",
            "Total",
            "Approved",
            "Rejected",
            "Quality Rate",
            "Risk",
        ]

        for index, header in enumerate(
            headers
        ):
            table.rows[0].cells[
                index
            ].text = header

        for row in context[
            "chemical_performance"
        ]:
            cells = table.add_row().cells

            values = [
                row.get("chemical"),
                row.get("total"),
                row.get("approved"),
                row.get("rejected"),
                row.get("quality_rate"),
                row.get("risk_score"),
            ]

            for index, value in enumerate(
                values
            ):
                cells[index].text = str(
                    value
                )

        document.add_heading(
            "Active Analysis",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=5,
        )

        headers = [
            "Process",
            "Chemical",
            "Status",
            "Age Hours",
            "Risk",
        ]

        for index, header in enumerate(
            headers
        ):
            table.rows[0].cells[
                index
            ].text = header

        for row in context[
            "active_processes"
        ][:100]:
            cells = table.add_row().cells

            values = [
                row.get("id"),
                row.get("chemical"),
                row.get("status"),
                row.get("age_hours"),
                row.get("risk"),
            ]

            for index, value in enumerate(
                values
            ):
                cells[index].text = str(
                    value
                )

        buffer = io.BytesIO()

        document.save(buffer)

        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
        )

        filename = self.get_export_file_name(
            "docx"
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; filename="{filename}"'
        )

        return response

    def export_dashboard_pdf(self):
        context = self.get_dashboard_context()

        html = """
        <html>
        <head>
        <style>
        body {
            font-family: Helvetica;
            font-size: 9px;
        }
        h1 {
            font-size: 18px;
        }
        h2 {
            font-size: 13px;
            margin-top: 16px;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 12px;
        }
        th, td {
            border: 0.5px solid #777;
            padding: 4px;
        }
        th {
            font-weight: bold;
        }
        </style>
        </head>
        <body>
        """

        html += (
            "<h1>"
            "Chemical Analysis Command Center"
            "</h1>"
        )

        html += (
            "<p>Reporting period: "
            f"{context['filters']['start']}"
            " to "
            f"{context['filters']['end']}"
            "</p>"
        )

        html += (
            "<h2>Executive KPIs</h2>"
            "<table>"
            "<tr><th>Metric</th><th>Value</th></tr>"
        )

        for key, value in context[
            "kpis"
        ].items():
            html += (
                "<tr>"
                f"<td>{key}</td>"
                f"<td>{value}</td>"
                "</tr>"
            )

        html += "</table>"

        html += (
            "<h2>Chemical Performance</h2>"
            "<table>"
            "<tr>"
            "<th>Chemical</th>"
            "<th>Total</th>"
            "<th>Approved</th>"
            "<th>Rejected</th>"
            "<th>Quality %</th>"
            "<th>Risk</th>"
            "</tr>"
        )

        for row in context[
            "chemical_performance"
        ][:50]:
            html += (
                "<tr>"
                f"<td>{row.get('chemical')}</td>"
                f"<td>{row.get('total')}</td>"
                f"<td>{row.get('approved')}</td>"
                f"<td>{row.get('rejected')}</td>"
                f"<td>{row.get('quality_rate')}</td>"
                f"<td>{row.get('risk_score')}</td>"
                "</tr>"
            )

        html += "</table>"

        html += (
            "<h2>Active Processes</h2>"
            "<table>"
            "<tr>"
            "<th>Process</th>"
            "<th>Chemical</th>"
            "<th>Status</th>"
            "<th>Age</th>"
            "<th>Risk</th>"
            "</tr>"
        )

        for row in context[
            "active_processes"
        ][:100]:
            html += (
                "<tr>"
                f"<td>{row.get('id')}</td>"
                f"<td>{row.get('chemical')}</td>"
                f"<td>{row.get('status')}</td>"
                f"<td>{row.get('age_hours')}</td>"
                f"<td>{row.get('risk')}</td>"
                "</tr>"
            )

        html += (
            "</table>"
            "</body>"
            "</html>"
        )

        response = HttpResponse(
            content_type="application/pdf"
        )

        filename = self.get_export_file_name(
            "pdf"
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; filename="{filename}"'
        )

        pisa.CreatePDF(
            io.BytesIO(
                html.encode("utf-8")
            ),
            dest=response,
        )

        return response

    def handle_export(self, request):
        export_type = (
            request.GET.get("export")
            or request.GET.get("type")
        )

        if not export_type:
            return None

        export_type = export_type.lower()

        export_map = {
            "csv": self.export_dashboard_csv,
            "xlsx": self.export_dashboard_excel,
            "excel": self.export_dashboard_excel,
            "docx": self.export_dashboard_docx,
            "word": self.export_dashboard_docx,
            "pdf": self.export_dashboard_pdf,
        }

        handler = export_map.get(
            export_type
        )

        if handler:
            return handler()

        if export_type in {
            "json",
            "dashboard-json",
            "analytics",
        }:
            return self.dashboard_json_response()

        return HttpResponse(
            "Unsupported dashboard export format.",
            status=400,
        )