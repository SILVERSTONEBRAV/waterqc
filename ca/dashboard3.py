from __future__ import annotations
import csv
import io
import json
import math
import statistics
import zipfile
from collections import Counter, defaultdict
from datetime import timedelta
from decimal import Decimal

import numpy as np
import pandas as pd
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.urls import NoReverseMatch, reverse
from django.utils import timezone
from django.views.generic import TemplateView

from .mixins import CaWanjalaMixin
from .navigation import CA_DASHBOARD_DATA_PERMISSIONS
from t.models import TreatmentPlant

from .models import (
    BatchTestResult,
    Chemical,
    ChemicalAnalysisInput,
    ChemicalAnalysisProcess,
    ChemicalBatch,
    ChemicalCoA,
    ChemicalAnalysisRequest,
    Supplier,

)

try:
    from scipy import stats as scipy_stats
except Exception:
    scipy_stats = None

try:
    from sklearn.ensemble import IsolationForest
except Exception:
    IsolationForest = None

try:
    from statsmodels.tsa.holtwinters import ExponentialSmoothing
except Exception:
    ExponentialSmoothing = None


User = get_user_model()


class Dashboard3View(CaWanjalaMixin, TemplateView):
    template_name = "ca/dashboard/pages/dashboard3.html"
    permission_required_view = "ca.view_chemicalanalysisprocess"

    process_statuses = (
        "CREATED",
        "RECEIVED",
        "REQUESTED",
        "CHECKLIST",
        "ASSIGNED",
        "SYSTEMATIC_RANDOM_SAMPLING",
        "ANALYSIS",
        "RETEST",
        "QA",
        "CAPA",
        "APPROVED",
        "REJECTED",
        "UTILIZED",
        "RETURNED",
        "COMPLETED",
    )

    decision_statuses = (
        "APPROVED",
        "REJECTED",
        "UTILIZED",
        "RETURNED",
        "COMPLETED",
    )

    result_statuses = ("PASS", "FAIL", "NA")

    period_map = {
        "7": 7,
        "30": 30,
        "90": 90,
        "365": 365,
    }

    group_map = {
        "day": "day",
        "week": "week",
        "month": "month",
    }

    compare_map = {
        "chemical": "chemical",
        "supplier": "supplier",
        "analyst": "analyst",
        "plant": "plant",
    }

    cache_timeout = 60

    def _get_model(self, name):
        from django.apps import apps
        try:
            return apps.get_model("ca", name)
        except LookupError:
            return None

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and not self.is_ca_superuser(request.user):
            if not all(
                request.user.has_perm(permission)
                for permission in CA_DASHBOARD_DATA_PERMISSIONS
            ):
                raise PermissionDenied(
                    "You do not have permission to view the full CA analytics dashboard."
                )
        self._now = timezone.now()
        self._filters = self._get_filters()
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        qs = (
            ChemicalAnalysisProcess.objects
            .select_related(
                "batch",
                "batch__chemical",
                "batch__supplier",
                "batch__treatment_plant",
                "analysis_request",
                "assigned_analyst",
                "qa_officer",
                "approver",
            )
            .prefetch_related(
                "batch__test_results",
                "batch__test_results__specification",
                "task_set",
            )
        )

        qs = self._scope_queryset(qs)
        qs = self._apply_filters(qs)

        return qs

    def _scope_queryset(self, qs):
        return self.scope_queryset(qs, field="batch__treatment_plant_id")

    def _get_filters(self):
        params = self.request.GET

        period = params.get("period", "30")
        if period not in self.period_map:
            period = "30"

        group = params.get("group", "day")
        if group not in self.group_map:
            group = "day"

        compare = params.get("compare", "chemical")
        if compare not in self.compare_map:
            compare = "chemical"

        return {
            "period": period,
            "days": self.period_map[period],
            "group": group,
            "compare": compare,
            "status": params.get("status", "").strip(),
            "chemical": params.get("chemical", "").strip(),
            "supplier": params.get("supplier", "").strip(),
            "analyst": params.get("analyst", "").strip(),
            "plant": params.get("plant", "").strip(),
            "date_from": params.get("date_from", "").strip(),
            "date_to": params.get("date_to", "").strip(),
            "min_turnaround": params.get("min_turnaround", "").strip(),
            "max_turnaround": params.get("max_turnaround", "").strip(),
            "q": params.get("q", "").strip(),
        }

    def _apply_filters(self, qs):
        f = self._filters

        if f["status"]:
            qs = qs.filter(status=f["status"])

        if f["chemical"]:
            qs = qs.filter(
                Q(batch__chemical__name__icontains=f["chemical"])
                | Q(batch__chemical__code__icontains=f["chemical"])
            )

        if f["supplier"]:
            qs = qs.filter(
                Q(batch__supplier__name__icontains=f["supplier"])
                | Q(batch__supplier__contract_number__icontains=f["supplier"])
            )

        if f["analyst"]:
            qs = qs.filter(
                Q(assigned_analyst__username__icontains=f["analyst"])
                | Q(assigned_analyst__first_name__icontains=f["analyst"])
                | Q(assigned_analyst__last_name__icontains=f["analyst"])
            )

        if f["plant"]:
            qs = qs.filter(batch__treatment_plant__name__icontains=f["plant"])

        if f["q"]:
            term = f["q"]
            qs = qs.filter(
                Q(batch__sample_id__icontains=term)
                | Q(batch__chemical__name__icontains=term)
                | Q(batch__chemical__code__icontains=term)
            )

        if f["date_from"]:
            try:
                dt = timezone.make_aware(
                    timezone.datetime.fromisoformat(f["date_from"])
                )
                qs = qs.filter(created_at__gte=dt)
            except ValueError:
                pass

        if f["date_to"]:
            try:
                dt = timezone.make_aware(
                    timezone.datetime.fromisoformat(f["date_to"])
                )
                qs = qs.filter(created_at__lt=dt + timedelta(days=1))
            except ValueError:
                pass

        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        processes = self.get_queryset()

        payload = self._build_dashboard(processes)

        context.update(payload)

        context["dashboard_url"] = self.request.path
        context["filters"] = self._filters
        context["available_periods"] = self.period_map.keys()
        context["permissions"] = {
            "export": self.user_has_ca_permission(
                "ca.view_chemicalanalysisprocess"
            ),
        }

        context["urls"] = {
            "dashboard": self.request.path,
            "processes": "/ca/processes/",
            "batches": "/ca/batches/",
            "requests": "/ca/requests/",
            "results": "/ca/results/",
            "coa": "/ca/coa/",
            "export": self.request.path,
        }

        return context

    def _build_dashboard(self, processes):
        if not hasattr(self, "_now") or self._now is None:
            self._now = timezone.now()
        if not hasattr(self, "_filters") or self._filters is None:
            self._filters = self._get_filters()

        cache_key = self._cache_key()

        if not self.request.GET.get("export"):
            cached = cache.get(cache_key)
            if cached:
                return cached

        process_list = list(processes)

        payload = {
            "meta": self._meta(process_list),
            "kpis": self._kpis(process_list),
            "workflow": self._workflow(process_list),
            "queues": self._queues(process_list),
            "quality": self._quality(process_list),
            "sla": self._sla(process_list),
            "turnaround": self._turnaround(process_list),
            "sampling": self._sampling(process_list),
            "comparatives": self._comparatives(process_list),
            "statistics": self._statistics(process_list),
            "risk": self._risk(process_list),
            "anomalies": self._anomalies(process_list),
            "drilldowns": self._drilldowns(process_list),
            "figures": self._figures(process_list),
            "charts": {},
        }

        payload["charts"] = payload["figures"]

        # -----------------------------------------------------------
        # Derived top-level keys expected by the template
        # -----------------------------------------------------------

        # risk_register – elevated-risk processes
        payload["risk_register"] = self._build_risk_register(process_list)

        # qa_queue / approval_queue – processes in QA/CAPA or APPROVED
        payload["qa_queue"] = self._build_queue(process_list, {"QA", "CAPA"})
        payload["approval_queue"] = self._build_queue(process_list, {"APPROVED"})

        # active_processes – all non-terminal processes
        terminal = {"COMPLETED", "REJECTED", "RETURNED", "UTILIZED"}
        payload["active_processes"] = self._build_active_processes(
            process_list, terminal,
        )

        # supplier_performance / analyst_performance tables
        payload["supplier_performance"] = self._build_supplier_performance(
            process_list,
        )
        payload["analyst_performance"] = self._build_analyst_performance(
            process_list,
        )

        # forecast availability flag
        payload["forecast"] = {"available": bool(process_list)}

        # comparisons (period-over-period deltas) – placeholder
        payload["comparisons"] = {}

        # comparative_statistics – placeholder so template |default works
        payload["comparative_statistics"] = {}

        if not self.request.GET.get("export"):
            cache.set(cache_key, payload, self.cache_timeout)

        return payload

    # -----------------------------------------------------------------
    # Helpers for derived template context
    # -----------------------------------------------------------------

    def _build_risk_register(self, processes):
        rows = []
        for p in processes:
            age = self._age_hours(p)
            rework = getattr(p, "rework_count", 0) or 0
            overdue = self._is_overdue(p)

            risk_score = 0
            if overdue:
                risk_score += 40
            if rework > 0:
                risk_score += 20 * rework
            if age > 48:
                risk_score += min(int(age / 12), 40)

            if risk_score <= 0:
                continue

            rows.append({
                "id": p.pk,
                "chemical": self._chemical_name(p),
                "status": p.status,
                "age_hours": round(age, 1),
                "rework_count": rework,
                "escalation_count": 0,
                "risk_score": risk_score,
            })

        rows.sort(key=lambda r: r["risk_score"], reverse=True)
        return rows

    def _build_queue(self, processes, statuses):
        rows = []
        for p in processes:
            if p.status not in statuses:
                continue
            age = self._age_hours(p)
            risk = "HIGH" if age > 24 else ("MEDIUM" if age > 12 else "LOW")
            rows.append({
                "id": p.pk,
                "status": p.status,
                "age_hours": round(age, 1),
                "date": getattr(p, "created_at", None),
                "risk": risk,
            })
        return rows

    def _build_active_processes(self, processes, terminal):
        rows = []
        active_task_statuses = {"NEW", "ASSIGNED", "STARTED"}
        for p in processes:
            if p.status in terminal:
                continue
            age = self._age_hours(p)
            overdue = self._is_overdue(p)
            rework = getattr(p, "rework_count", 0) or 0

            if age > 48 or overdue:
                risk = "CRITICAL" if (overdue and rework > 0) else "HIGH"
            elif age > 24:
                risk = "MEDIUM"
            else:
                risk = "LOW"

            action_url = reverse("ca:process-detail", kwargs={"pk": p.pk})
            action_title = "Open process"
            for task in sorted(
                (
                    task
                    for task in p.task_set.all()
                    if task.status in active_task_statuses and task.flow_task
                ),
                key=lambda task: (task.created, task.pk),
            ):
                node = task.flow_task
                try:
                    available = (
                        node.can_execute(self.request.user, task)
                        or node.can_assign(self.request.user, task)
                        or node.can_view(self.request.user, task)
                    )
                except AttributeError:
                    available = False
                if available:
                    try:
                        action_url = reverse(
                            f"ca:chemicalanalysis:{node.name}:index",
                            args=[p.pk, task.pk],
                        )
                    except (NoReverseMatch, AttributeError, TypeError):
                        # Fall back to the process page rather than failing the
                        # entire dashboard on a bad Viewflow namespace.
                        pass
                    action_title = task.title or str(node)
                    break

            rows.append({
                "id": p.pk,
                "chemical": self._chemical_name(p),
                "status": p.status,
                "age_hours": round(age, 1),
                "rework_count": rework,
                "overdue": overdue,
                "risk": risk,
                "action_url": action_url,
                "action_title": action_title,
            })
        return rows

    def _build_supplier_performance(self, processes):
        groups = defaultdict(lambda: Counter())
        for p in processes:
            name = self._supplier_name(p)
            groups[name][p.status] += 1

        rows = []
        for supplier, statuses in groups.items():
            total = sum(statuses.values())
            approved = statuses.get("APPROVED", 0) + statuses.get("COMPLETED", 0) + statuses.get("UTILIZED", 0)
            rejected = statuses.get("REJECTED", 0) + statuses.get("RETURNED", 0)
            acceptance = round(approved / total * 100, 1) if total else 0
            risk_score = round(100 - acceptance, 1)
            rows.append({
                "supplier": supplier,
                "total": total,
                "approved": approved,
                "rejected": rejected,
                "acceptance_rate": acceptance,
                "risk_score": risk_score,
            })
        rows.sort(key=lambda r: r["total"], reverse=True)
        return rows

    def _build_analyst_performance(self, processes):
        groups = defaultdict(lambda: Counter())
        terminal = {"COMPLETED", "REJECTED", "RETURNED", "UTILIZED"}
        for p in processes:
            name = self._user_name(getattr(p, "assigned_analyst", None))
            groups[name][p.status] += 1

        rows = []
        for analyst, statuses in groups.items():
            total = sum(statuses.values())
            active = sum(c for s, c in statuses.items() if s not in terminal)
            completed = statuses.get("COMPLETED", 0)
            retest = statuses.get("RETEST", 0)
            qa = statuses.get("QA", 0) + statuses.get("CAPA", 0)
            rows.append({
                "analyst": analyst,
                "total": total,
                "active": active,
                "completed": completed,
                "retest": retest,
                "qa": qa,
                "capa": statuses.get("CAPA", 0),
            })
        rows.sort(key=lambda r: r["total"], reverse=True)
        return rows

    def _age_hours(self, process):
        created = getattr(process, "created_at", None)
        if created is None:
            return 0
        return max((self._now - created).total_seconds() / 3600, 0)

    def _cache_key(self):
        user = self.request.user
        plant_id = self.get_scope_plant_id()
        scope = (
            "all"
            if plant_id is None and self.is_ca_superuser(user)
            else plant_id
        )

        query = self.request.GET.urlencode()

        # The payload can contain plant-scoped process data. Keep entries
        # isolated by user and the effective selected plant, including the
        # superuser plant selector stored in the session.
        return f"ca-dashboard3:{user.pk}:{scope}:{query}"

    def _meta(self, processes):
        return {
            "generated_at": self._now.isoformat(),
            "timezone": timezone.get_current_timezone_name(),
            "period": self._filters["period"],
            "days": self._filters["days"],
            "group": self._filters["group"],
            "compare": self._filters["compare"],
            "process_count": len(processes),
            "scope": (
                "global"
                if self.request.user.is_superuser
                else str(
                    getattr(
                        getattr(
                            self.request.user,
                            "assigned_treatment_plant",
                            None,
                        ),
                        "name",
                        "Unassigned",
                    )
                )
            ),
        }

    def _kpis(self, processes):
        total = len(processes)

        active = sum(
            1
            for p in processes
            if p.status not in {"COMPLETED", "REJECTED", "RETURNED", "UTILIZED"}
        )

        completed = sum(1 for p in processes if p.status == "COMPLETED")

        overdue = sum(1 for p in processes if self._is_overdue(p))

        results = self._results(processes)
        pass_count = results["PASS"]
        fail_count = results["FAIL"]
        na_count = results["NA"]
        result_total = sum(results.values())

        turnaround = self._turnaround_values(processes)

        retested = sum(
            1
            for p in processes
            if getattr(p, "rework_count", 0) and p.rework_count > 0
        )

        analysed = sum(
            1
            for p in processes
            if getattr(p, "analysis_started_at", None)
        )

        sla_rate = (
            ((total - overdue) / total) * 100
            if total
            else 0
        )

        pass_rate = (
            pass_count / result_total * 100
            if result_total
            else 0
        )

        retest_rate = (
            retested / analysed * 100
            if analysed
            else 0
        )

        fail_rate = (
            fail_count / result_total * 100
            if result_total
            else 0
        )

        return {
            "total_processes": total,
            "total_batches": len({p.batch_id for p in processes}),
            "active_processes": active,
            "active_analyses": active,
            "completed_processes": completed,
            "completed": completed,
            "overdue_processes": overdue,
            "pass_count": pass_count,
            "fail_count": fail_count,
            "na_count": na_count,
            "pass_rate": round(pass_rate, 2),
            "fail_rate": round(fail_rate, 2),
            "retest_rate": round(retest_rate, 2),
            "retest_count": retested,
            "sla_rate": round(sla_rate, 2),
            "avg_turnaround": self._mean(turnaround),
            "median_turnaround": self._median(turnaround),
            "p90_turnaround": self._percentile(turnaround, 90),
            "p95_turnaround": self._percentile(turnaround, 95),
        }

    def _workflow(self, processes):
        counts = Counter(p.status for p in processes)

        stages = {}

        for status in self.process_statuses:
            stages[status.lower()] = {
                "code": status,
                "label": status.replace("_", " ").title(),
                "count": counts.get(status, 0),
            }

        conversions = []

        ordered = [
            "CREATED",
            "RECEIVED",
            "REQUESTED",
            "CHECKLIST",
            "ASSIGNED",
            "SYSTEMATIC_RANDOM_SAMPLING",
            "ANALYSIS",
            "QA",
            "APPROVED",
            "COMPLETED",
        ]

        for previous, current in zip(ordered, ordered[1:]):
            previous_count = counts.get(previous, 0)
            current_count = counts.get(current, 0)

            conversions.append(
                {
                    "from": previous,
                    "to": current,
                    "from_count": previous_count,
                    "to_count": current_count,
                    "conversion_rate": (
                        round(current_count / previous_count * 100, 2)
                        if previous_count
                        else 0
                    ),
                }
            )

        return {
            "stages": stages,
            "conversion": conversions,
        }

    def _queues(self, processes):
        SamplingModel = self._get_model("SystematicRandomSampling")
        if SamplingModel:
            sampling = list(
                SamplingModel.objects
                .filter(batch__in=[p.batch_id for p in processes])
            )
        else:
            sampling = []

        sampling_pending = sum(
            1 for s in sampling if getattr(s, "status", None) == "PENDING"
        )

        counter_check = sum(
            1 for s in sampling if getattr(s, "status", None) == "COUNTER_CHECK"
        )

        sampling_qa = sum(
            1 for s in sampling if getattr(s, "status", None) == "QA_REVIEW"
        )

        pending_results = 0
        method_inputs = 0
        qa = 0
        approval = 0
        coa = 0

        for process in processes:
            results = list(process.batch.test_results.all())

            pending_results += sum(
                1
                for result in results
                if not getattr(result, "status", None)
            )

            method_inputs += sum(
                1
                for result in results
                if not hasattr(result, "inputs")
            )

            if process.status in {"QA", "CAPA"}:
                qa += 1

            if process.status == "APPROVED":
                approval += 1

            if process.status in {"APPROVED", "COMPLETED"}:
                try:
                    process.batch.coa
                    coa += 1
                except Exception:
                    pass

        return {
            "sampling_pending": sampling_pending,
            "counter_check": counter_check,
            "sampling_qa": sampling_qa,
            "pending_results": pending_results,
            "method_inputs": method_inputs,
            "qa": qa,
            "approval": approval,
            "coa": coa,
        }

    def _quality(self, processes):
        result_counts = self._results(processes)

        chemical = defaultdict(lambda: Counter())
        supplier = defaultdict(lambda: Counter())
        analyst = defaultdict(lambda: Counter())

        for process in processes:
            chemical_name = self._chemical_name(process)
            supplier_name = self._supplier_name(process)
            analyst_name = self._user_name(
                getattr(process, "assigned_analyst", None)
            )

            for result in process.batch.test_results.all():
                status = getattr(result, "status", None)

                if status:
                    chemical[chemical_name][status] += 1
                    supplier[supplier_name][status] += 1
                    analyst[analyst_name][status] += 1

        retested = sum(
            1
            for p in processes
            if getattr(p, "rework_count", 0) and p.rework_count > 0
        )

        return {
            "results": dict(result_counts),
            "pass": result_counts.get("PASS", 0),
            "fail": result_counts.get("FAIL", 0),
            "na": result_counts.get("NA", 0),
            "retest_count": retested,
            "chemical": self._counter_dict(chemical),
            "supplier": self._counter_dict(supplier),
            "analyst": self._counter_dict(analyst),
        }

    def _sla(self, processes):
        buckets = {
            "future": 0,
            "due_soon": 0,
            "overdue_0_3h": 0,
            "overdue_3_6h": 0,
            "overdue_6_12h": 0,
            "overdue_12_24h": 0,
            "overdue_24h_plus": 0,
        }

        for process in processes:
            deadline = self._deadline(process)

            if deadline is None:
                continue

            delta = (self._now - deadline).total_seconds() / 3600

            if delta < -3:
                buckets["future"] += 1
            elif delta < 0:
                buckets["due_soon"] += 1
            elif delta < 3:
                buckets["overdue_0_3h"] += 1
            elif delta < 6:
                buckets["overdue_3_6h"] += 1
            elif delta < 12:
                buckets["overdue_6_12h"] += 1
            elif delta < 24:
                buckets["overdue_12_24h"] += 1
            else:
                buckets["overdue_24h_plus"] += 1

        total = len(processes)
        overdue = sum(
            buckets[key]
            for key in (
                "overdue_0_3h",
                "overdue_3_6h",
                "overdue_6_12h",
                "overdue_12_24h",
                "overdue_24h_plus",
            )
        )

        return {
            "buckets": buckets,
            "total": total,
            "overdue": overdue,
            "compliant": max(total - overdue, 0),
            "rate": round(
                ((total - overdue) / total) * 100,
                2,
            ) if total else 0,
        }

    def _turnaround(self, processes):
        values = self._turnaround_values(processes)

        if not values:
            return {
                "count": 0,
                "mean": 0,
                "median": 0,
                "min": 0,
                "max": 0,
                "p25": 0,
                "p75": 0,
                "p90": 0,
                "p95": 0,
                "std": 0,
                "iqr": 0,
                "cv": 0,
                "skew": 0,
            }

        mean = self._mean(values)
        std = self._std(values)

        return {
            "count": len(values),
            "mean": mean,
            "median": self._median(values),
            "min": min(values),
            "max": max(values),
            "p25": self._percentile(values, 25),
            "p75": self._percentile(values, 75),
            "p90": self._percentile(values, 90),
            "p95": self._percentile(values, 95),
            "std": std,
            "iqr": round(
                self._percentile(values, 75)
                - self._percentile(values, 25),
                2,
            ),
            "cv": round(
                std / mean * 100,
                2,
            ) if mean else 0,
            "skew": self._skew(values),
        }

    def _sampling(self, processes):
        batch_ids = [p.batch_id for p in processes]

        SamplingModel = self._get_model("SystematicRandomSampling")
        if SamplingModel:
            records = list(
                SamplingModel.objects
                .filter(batch_id__in=batch_ids)
                .prefetch_related("samples")
            )
        else:
            records = []

        status_counts = Counter(getattr(r, "status", "") for r in records)

        package_variances = []

        for record in records:
            for sample in record.samples.all():
                declared = self._decimal(
                    getattr(sample, "declared_weight", None)
                )
                measured = self._decimal(
                    getattr(sample, "measured_weight", None)
                )

                if declared is None or measured is None:
                    continue

                package_variances.append(
                    float(measured - declared)
                )

        return {
            "records": len(records),
            "status": dict(status_counts),
            "samples": sum(
                getattr(r, "samples").count()
                for r in records
            ),
            "weight_variance": {
                "mean": self._mean(package_variances),
                "median": self._median(package_variances),
                "max_abs": (
                    max(abs(x) for x in package_variances)
                    if package_variances
                    else 0
                ),
            },
        }

    def _comparatives(self, processes):
        dimensions = {
            "chemical": defaultdict(
                lambda: Counter()
            ),
            "supplier": defaultdict(
                lambda: Counter()
            ),
            "analyst": defaultdict(
                lambda: Counter()
            ),
            "plant": defaultdict(
                lambda: Counter()
            ),
        }

        turnaround = defaultdict(list)

        for process in processes:
            keys = {
                "chemical": self._chemical_name(process),
                "supplier": self._supplier_name(process),
                "analyst": self._user_name(
                    getattr(process, "assigned_analyst", None)
                ),
                "plant": getattr(
                    getattr(
                        process.batch,
                        "treatment_plant",
                        None,
                    ),
                    "name",
                    "Unassigned",
                ),
            }

            for dimension, key in keys.items():
                dimensions[dimension][key][process.status] += 1

                value = self._turnaround_for_process(process)
                if value is not None:
                    turnaround[(dimension, key)].append(value)

        output = {}

        for dimension, groups in dimensions.items():
            rows = []

            for key, statuses in groups.items():
                values = turnaround[(dimension, key)]
                total = sum(statuses.values())

                rows.append(
                    {
                        "name": key,
                        "total": total,
                        "completed": statuses.get("COMPLETED", 0),
                        "approved": statuses.get("APPROVED", 0),
                        "rejected": statuses.get("REJECTED", 0),
                        "returned": statuses.get("RETURNED", 0),
                        "active": sum(
                            count
                            for status, count in statuses.items()
                            if status not in {
                                "COMPLETED",
                                "REJECTED",
                                "RETURNED",
                                "UTILIZED",
                            }
                        ),
                        "turnaround": self._mean(values),
                    }
                )

            rows.sort(
                key=lambda item: item["total"],
                reverse=True,
            )

            output[dimension] = rows

        return output

    def _statistics(self, processes):
        values = self._turnaround_values(processes)

        if not values:
            return {}

        array = np.asarray(values, dtype=float)

        stats = {
            "count": int(array.size),
            "mean": round(float(np.mean(array)), 2),
            "median": round(float(np.median(array)), 2),
            "std": round(float(np.std(array)), 2),
            "min": round(float(np.min(array)), 2),
            "max": round(float(np.max(array)), 2),
            "p25": round(float(np.percentile(array, 25)), 2),
            "p75": round(float(np.percentile(array, 75)), 2),
            "p90": round(float(np.percentile(array, 90)), 2),
            "p95": round(float(np.percentile(array, 95)), 2),
            "iqr": round(
                float(
                    np.percentile(array, 75)
                    - np.percentile(array, 25)
                ),
                2,
            ),
        }

        if stats["mean"]:
            stats["cv"] = round(
                stats["std"] / stats["mean"] * 100,
                2,
            )
        else:
            stats["cv"] = 0

        if scipy_stats and len(array) >= 3:
            try:
                ci = scipy_stats.t.interval(
                    0.95,
                    len(array) - 1,
                    loc=np.mean(array),
                    scale=scipy_stats.sem(array),
                )
                stats["confidence_interval_95"] = [
                    round(float(ci[0]), 2),
                    round(float(ci[1]), 2),
                ]
            except Exception:
                stats["confidence_interval_95"] = None

        stats["control_limits"] = self._control_limits(array)

        return stats

    def _risk(self, processes):
        risks = []

        overdue = [p for p in processes if self._is_overdue(p)]

        if overdue:
            risks.append(
                {
                    "code": "SLA",
                    "level": "HIGH" if len(overdue) >= 5 else "MEDIUM",
                    "title": "SLA exposure",
                    "count": len(overdue),
                    "description": (
                        f"{len(overdue)} processes are beyond "
                        "their current deadline."
                    ),
                }
            )

        failed = self._results(processes)["FAIL"]

        if failed:
            risks.append(
                {
                    "code": "QUALITY",
                    "level": "HIGH" if failed >= 5 else "MEDIUM",
                    "title": "Failed analytical results",
                    "count": failed,
                    "description": (
                        f"{failed} analytical results are currently "
                        "recorded as FAIL."
                    ),
                }
            )

        retests = sum(
            1
            for p in processes
            if getattr(p, "rework_count", 0) > 0
        )

        if retests:
            risks.append(
                {
                    "code": "RETEST",
                    "level": "MEDIUM",
                    "title": "Retest activity",
                    "count": retests,
                    "description": (
                        f"{retests} processes have recorded rework."
                    ),
                }
            )

        sampling = self._sampling(processes)

        if sampling["weight_variance"]["max_abs"] > 0:
            risks.append(
                {
                    "code": "SAMPLING",
                    "level": "MEDIUM",
                    "title": "Sampling weight variance",
                    "count": sampling["samples"],
                    "description": (
                        "Measured package weights differ from "
                        "declared weights."
                    ),
                }
            )

        return risks

    def _anomalies(self, processes):
        values = []

        for process in processes:
            turnaround = self._turnaround_for_process(process)

            if turnaround is not None:
                values.append(
                    {
                        "process": process,
                        "value": turnaround,
                    }
                )

        if len(values) < 5:
            return []

        numbers = np.asarray(
            [item["value"] for item in values],
            dtype=float,
        )

        mean = float(np.mean(numbers))
        std = float(np.std(numbers))

        anomalies = []

        if std > 0:
            for item in values:
                z = (item["value"] - mean) / std

                if abs(z) >= 3:
                    anomalies.append(
                        {
                            "process_id": item["process"].pk,
                            "sample_id": getattr(
                                item["process"].batch,
                                "sample_id",
                                "",
                            ),
                            "value": round(item["value"], 2),
                            "z_score": round(z, 2),
                            "type": "turnaround",
                        }
                    )

        if IsolationForest and len(numbers) >= 10:
            try:
                model = IsolationForest(
                    contamination="auto",
                    random_state=42,
                )

                predictions = model.fit_predict(
                    numbers.reshape(-1, 1)
                )

                for index, prediction in enumerate(predictions):
                    if prediction == -1:
                        process = values[index]["process"]

                        if not any(
                            x["process_id"] == process.pk
                            for x in anomalies
                        ):
                            anomalies.append(
                                {
                                    "process_id": process.pk,
                                    "sample_id": getattr(
                                        process.batch,
                                        "sample_id",
                                        "",
                                    ),
                                    "value": round(
                                        values[index]["value"],
                                        2,
                                    ),
                                    "z_score": None,
                                    "type": "isolation_forest",
                                }
                            )
            except Exception:
                pass

        return anomalies[:100]

    def _drilldowns(self, processes):
        return {
            "status": self._status_drilldown(processes),
            "chemical": self._dimension_drilldown(
                processes,
                "chemical",
            ),
            "supplier": self._dimension_drilldown(
                processes,
                "supplier",
            ),
            "analyst": self._dimension_drilldown(
                processes,
                "analyst",
            ),
            "plant": self._dimension_drilldown(
                processes,
                "plant",
            ),
        }

    def _status_drilldown(self, processes):
        result = []

        groups = defaultdict(list)

        for process in processes:
            groups[process.status].append(process)

        for status, rows in groups.items():
            result.append(
                {
                    "status": status,
                    "count": len(rows),
                    "process_ids": [
                        p.pk for p in rows[:100]
                    ],
                    "sample_ids": [
                        getattr(p.batch, "sample_id", "")
                        for p in rows[:100]
                    ],
                }
            )

        return result

    def _dimension_drilldown(self, processes, dimension):
        groups = defaultdict(list)

        for process in processes:
            if dimension == "chemical":
                key = self._chemical_name(process)
            elif dimension == "supplier":
                key = self._supplier_name(process)
            elif dimension == "analyst":
                key = self._user_name(
                    getattr(process, "assigned_analyst", None)
                )
            else:
                key = getattr(
                    getattr(
                        process.batch,
                        "treatment_plant",
                        None,
                    ),
                    "name",
                    "Unassigned",
                )

            groups[key].append(process)

        return [
            {
                "name": key,
                "count": len(rows),
                "process_ids": [
                    p.pk for p in rows[:100]
                ],
            }
            for key, rows in sorted(
                groups.items(),
                key=lambda item: len(item[1]),
                reverse=True,
            )
        ]

    def _figures(self, processes):
        return {
            "throughput": self._plotly(
                self._throughput_figure(processes)
            ),
            "volume": self._plotly(
                self._throughput_figure(processes)
            ),
            "workflow": self._plotly(
                self._workflow_figure(processes)
            ),
            "quality": self._plotly(
                self._quality_figure(processes)
            ),
            "quality_trend": self._plotly(
                self._quality_trend_figure(processes)
            ),
            "turnaround": self._plotly(
                self._turnaround_figure(processes)
            ),
            "sla": self._plotly(
                self._sla_figure(processes)
            ),
            "sampling": self._plotly(
                self._sampling_figure(processes)
            ),
            "comparative": self._plotly(
                self._comparative_figure(processes)
            ),
            "chemical": self._plotly(
                self._dimension_figure(
                    processes,
                    "chemical",
                )
            ),
            "analyst": self._plotly(
                self._dimension_figure(
                    processes,
                    "analyst",
                )
            ),
            "supplier": self._plotly(
                self._dimension_figure(
                    processes,
                    "supplier",
                )
            ),
            "failures": self._plotly(
                self._pareto_figure(processes)
            ),
            "pareto": self._plotly(
                self._pareto_figure(processes)
            ),
            "correlation": self._plotly(
                self._correlation_figure(processes)
            ),
            "spc": self._plotly(
                self._spc_figure(processes)
            ),
            "forecast": self._plotly(
                self._forecast_figure(processes)
            ),
            "bottleneck": self._plotly(
                self._bottleneck_figure(processes)
            ),
            "heatmap": self._plotly(
                self._heatmap_figure(processes)
            ),
        }

    def _throughput_figure(self, processes):
        daily = defaultdict(
            lambda: {
                "created": 0,
                "completed": 0,
                "failed": 0,
                "overdue": 0,
            }
        )

        start = self._now - timedelta(
            days=self._filters["days"]
        )

        for process in processes:
            created = getattr(process, "created_at", None)

            if created and created >= start:
                key = created.date().isoformat()
                daily[key]["created"] += 1

            completed = getattr(
                process,
                "completed_at",
                None,
            )

            if completed and completed >= start:
                key = completed.date().isoformat()
                daily[key]["completed"] += 1

            if process.status == "REJECTED":
                if created and created >= start:
                    daily[created.date().isoformat()]["failed"] += 1

            if self._is_overdue(process):
                if created and created >= start:
                    daily[created.date().isoformat()]["overdue"] += 1

        dates = sorted(daily)

        return {
            "data": [
                {
                    "type": "bar",
                    "name": "Created",
                    "x": dates,
                    "y": [
                        daily[d]["created"]
                        for d in dates
                    ],
                },
                {
                    "type": "bar",
                    "name": "Completed",
                    "x": dates,
                    "y": [
                        daily[d]["completed"]
                        for d in dates
                    ],
                },
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": "Failed",
                    "x": dates,
                    "y": [
                        daily[d]["failed"]
                        for d in dates
                    ],
                },
                {
                    "type": "scatter",
                    "mode": "lines",
                    "name": "Overdue",
                    "x": dates,
                    "y": [
                        daily[d]["overdue"]
                        for d in dates
                    ],
                },
            ],
            "layout": self._layout(
                "Throughput and Process Load"
            ),
        }

    def _workflow_figure(self, processes):
        counts = Counter(
            p.status for p in processes
        )

        labels = [
            status.replace("_", " ").title()
            for status in self.process_statuses
        ]

        values = [
            counts.get(status, 0)
            for status in self.process_statuses
        ]

        return {
            "data": [
                {
                    "type": "funnel",
                    "y": labels,
                    "x": values,
                    "textinfo": "value+percent initial",
                }
            ],
            "layout": self._layout(
                "Workflow Funnel"
            ),
        }

    def _quality_figure(self, processes):
        results = self._results(processes)

        return {
            "data": [
                {
                    "type": "pie",
                    "labels": [
                        "PASS",
                        "FAIL",
                        "NA",
                    ],
                    "values": [
                        results["PASS"],
                        results["FAIL"],
                        results["NA"],
                    ],
                    "hole": 0.55,
                    "textinfo": "label+percent",
                }
            ],
            "layout": self._layout(
                "Analytical Result Distribution"
            ),
        }

    def _quality_trend_figure(self, processes):
        daily = defaultdict(
            lambda: Counter()
        )

        for process in processes:
            created = getattr(process, "created_at", None)

            if not created:
                continue

            date_key = created.date().isoformat()

            for result in process.batch.test_results.all():
                status = getattr(
                    result,
                    "status",
                    None,
                )

                if status:
                    daily[date_key][status] += 1

        dates = sorted(daily)

        return {
            "data": [
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": status,
                    "x": dates,
                    "y": [
                        daily[d][status]
                        for d in dates
                    ],
                }
                for status in self.result_statuses
            ],
            "layout": self._layout(
                "Quality Trend"
            ),
        }

    def _turnaround_figure(self, processes):
        values = self._turnaround_values(processes)

        return {
            "data": [
                {
                    "type": "box",
                    "y": values,
                    "name": "Turnaround",
                    "boxpoints": "outliers",
                }
            ],
            "layout": self._layout(
                "Turnaround Distribution",
                yaxis_title="Hours",
            ),
        }

    def _sla_figure(self, processes):
        sla = self._sla(processes)
        buckets = sla["buckets"]

        labels = list(buckets)
        values = list(buckets.values())

        return {
            "data": [
                {
                    "type": "bar",
                    "x": labels,
                    "y": values,
                    "text": values,
                    "textposition": "auto",
                }
            ],
            "layout": self._layout(
                "SLA Aging"
            ),
        }

    def _sampling_figure(self, processes):
        sampling = self._sampling(processes)

        labels = list(
            sampling["status"].keys()
        )

        values = [
            sampling["status"][x]
            for x in labels
        ]

        return {
            "data": [
                {
                    "type": "bar",
                    "x": labels,
                    "y": values,
                }
            ],
            "layout": self._layout(
                "Sampling Queue"
            ),
        }

    def _comparative_figure(self, processes):
        dimension = self._filters["compare"]

        rows = self._comparatives(
            processes
        ).get(dimension, [])

        rows = rows[:20]

        return {
            "data": [
                {
                    "type": "bar",
                    "name": "Total",
                    "x": [
                        row["name"]
                        for row in rows
                    ],
                    "y": [
                        row["total"]
                        for row in rows
                    ],
                },
                {
                    "type": "bar",
                    "name": "Completed",
                    "x": [
                        row["name"]
                        for row in rows
                    ],
                    "y": [
                        row["completed"]
                        for row in rows
                    ],
                },
            ],
            "layout": self._layout(
                f"{dimension.title()} Comparison",
                barmode="group",
            ),
        }

    def _dimension_figure(self, processes, dimension):
        rows = self._comparatives(
            processes
        ).get(dimension, [])[:20]

        names = [
            row["name"]
            for row in rows
        ]

        total = [
            row["total"]
            for row in rows
        ]

        turnaround = [
            row["turnaround"]
            for row in rows
        ]

        return {
            "data": [
                {
                    "type": "bar",
                    "name": "Volume",
                    "x": names,
                    "y": total,
                },
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": "Mean Turnaround",
                    "x": names,
                    "y": turnaround,
                    "yaxis": "y2",
                },
            ],
            "layout": self._layout(
                f"{dimension.title()} Performance",
                yaxis2={
                    "title": "Hours",
                    "overlaying": "y",
                    "side": "right",
                },
            ),
        }

    def _pareto_figure(self, processes):
        failures = Counter()

        for process in processes:
            for result in process.batch.test_results.all():
                if getattr(result, "status", None) == "FAIL":
                    name = getattr(
                        getattr(
                            result,
                            "specification",
                            None,
                        ),
                        "parameter",
                        None,
                    )

                    name = (
                        getattr(name, "name", None)
                        or getattr(
                            getattr(
                                result,
                                "specification",
                                None,
                            ),
                            "name",
                            None,
                        )
                        or "Unclassified"
                    )

                    failures[name] += 1

        rows = failures.most_common(20)

        names = [
            item[0] for item in rows
        ]

        counts = [
            item[1] for item in rows
        ]

        cumulative = []
        total = sum(counts)
        running = 0

        for count in counts:
            running += count
            cumulative.append(
                running / total * 100
                if total
                else 0
            )

        return {
            "data": [
                {
                    "type": "bar",
                    "name": "Failures",
                    "x": names,
                    "y": counts,
                },
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": "Cumulative %",
                    "x": names,
                    "y": cumulative,
                    "yaxis": "y2",
                },
            ],
            "layout": self._layout(
                "Failure Pareto",
                yaxis2={
                    "title": "Cumulative %",
                    "overlaying": "y",
                    "side": "right",
                    "range": [0, 100],
                },
            ),
        }

    def _correlation_figure(self, processes):
        rows = []

        for process in processes:
            rows.append(
                {
                    "turnaround": self._turnaround_for_process(
                        process
                    ),
                    "rework": getattr(
                        process,
                        "rework_count",
                        0,
                    ),
                    "overdue": int(
                        self._is_overdue(process)
                    ),
                    "failed": sum(
                        1
                        for result in process.batch.test_results.all()
                        if getattr(
                            result,
                            "status",
                            None,
                        ) == "FAIL"
                    ),
                }
            )

        df = pd.DataFrame(rows)

        if df.empty:
            matrix = []
            labels = []
        else:
            df = df.infer_objects(copy=False)
            numeric_columns = df.select_dtypes(include=["number", "bool"]).columns
            df[numeric_columns] = df[numeric_columns].fillna(0)
            corr = df.corr(numeric_only=True)
            labels = list(corr.columns)
            matrix = corr.round(3).values.tolist()

        return {
            "data": [
                {
                    "type": "heatmap",
                    "z": matrix,
                    "x": labels,
                    "y": labels,
                    "zmin": -1,
                    "zmax": 1,
                    "text": matrix,
                    "texttemplate": "%{text}",
                }
            ],
            "layout": self._layout(
                "Process Metric Correlation"
            ),
        }

    def _spc_figure(self, processes):
        values = []

        for process in processes:
            value = self._turnaround_for_process(
                process
            )

            if value is not None:
                values.append(value)

        if not values:
            return {
                "data": [],
                "layout": self._layout(
                    "Turnaround Control Chart"
                ),
            }

        array = np.asarray(values)
        mean = float(np.mean(array))
        std = float(np.std(array))

        ucl = mean + 3 * std
        lcl = max(mean - 3 * std, 0)

        x = list(range(1, len(values) + 1))

        return {
            "data": [
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": "Turnaround",
                    "x": x,
                    "y": values,
                },
                {
                    "type": "scatter",
                    "mode": "lines",
                    "name": "Mean",
                    "x": x,
                    "y": [mean] * len(x),
                },
                {
                    "type": "scatter",
                    "mode": "lines",
                    "name": "UCL",
                    "x": x,
                    "y": [ucl] * len(x),
                },
                {
                    "type": "scatter",
                    "mode": "lines",
                    "name": "LCL",
                    "x": x,
                    "y": [lcl] * len(x),
                },
            ],
            "layout": self._layout(
                "Turnaround Statistical Process Control",
                yaxis_title="Hours",
            ),
        }

    def _forecast_figure(self, processes):
        daily = Counter()

        for process in processes:
            created = getattr(
                process,
                "created_at",
                None,
            )

            if created:
                daily[created.date()] += 1

        if len(daily) < 7:
            return {
                "data": [],
                "layout": self._layout(
                    "Throughput Forecast"
                ),
            }

        dates = sorted(daily)
        values = [
            daily[d]
            for d in dates
        ]

        forecast_days = 14

        future_dates = []
        forecast_values = []

        if ExponentialSmoothing:
            try:
                model = ExponentialSmoothing(
                    values,
                    trend="add",
                    seasonal=None,
                ).fit()

                prediction = model.forecast(
                    forecast_days
                )

                future_dates = [
                    dates[-1] + timedelta(days=i)
                    for i in range(1, forecast_days + 1)
                ]

                forecast_values = [
                    max(float(value), 0)
                    for value in prediction
                ]
            except Exception:
                pass

        if not forecast_values:
            mean = (
                sum(values[-7:]) / min(7, len(values))
            )

            future_dates = [
                dates[-1] + timedelta(days=i)
                for i in range(1, forecast_days + 1)
            ]

            forecast_values = [
                mean for _ in future_dates
            ]

        return {
            "data": [
                {
                    "type": "scatter",
                    "mode": "lines+markers",
                    "name": "Actual",
                    "x": [
                        d.isoformat()
                        for d in dates
                    ],
                    "y": values,
                },
                {
                    "type": "scatter",
                    "mode": "lines",
                    "name": "Forecast",
                    "x": [
                        d.isoformat()
                        for d in future_dates
                    ],
                    "y": forecast_values,
                },
            ],
            "layout": self._layout(
                "Throughput Forecast"
            ),
        }

    def _bottleneck_figure(self, processes):
        statuses = [
            "CREATED",
            "RECEIVED",
            "REQUESTED",
            "CHECKLIST",
            "ASSIGNED",
            "SYSTEMATIC_RANDOM_SAMPLING",
            "ANALYSIS",
            "QA",
            "APPROVED",
            "COMPLETED",
        ]

        counts = Counter(
            p.status for p in processes
        )

        gaps = []

        for first, second in zip(
            statuses,
            statuses[1:],
        ):
            first_count = counts.get(first, 0)
            second_count = counts.get(second, 0)

            gaps.append(
                max(first_count - second_count, 0)
            )

        return {
            "data": [
                {
                    "type": "bar",
                    "x": [
                        f"{a} → {b}"
                        for a, b in zip(
                            statuses,
                            statuses[1:],
                        )
                    ],
                    "y": gaps,
                }
            ],
            "layout": self._layout(
                "Workflow Bottleneck Gaps"
            ),
        }

    def _heatmap_figure(self, processes):
        matrix = defaultdict(
            lambda: Counter()
        )

        for process in processes:
            created = getattr(
                process,
                "created_at",
                None,
            )

            if not created:
                continue

            weekday = created.strftime("%A")
            hour = created.hour

            matrix[weekday][hour] += 1

        weekdays = [
            "Monday",
            "Tuesday",
            "Wednesday",
            "Thursday",
            "Friday",
            "Saturday",
            "Sunday",
        ]

        hours = list(range(24))

        z = [
            [
                matrix[weekday][hour]
                for hour in hours
            ]
            for weekday in weekdays
        ]

        return {
            "data": [
                {
                    "type": "heatmap",
                    "x": hours,
                    "y": weekdays,
                    "z": z,
                }
            ],
            "layout": self._layout(
                "Process Creation Heatmap"
            ),
        }

    def _plotly(self, figure):
        return figure or {
            "data": [],
            "layout": self._layout(
                "No data"
            ),
        }

    def _layout(self, title, **kwargs):
        layout = {
            "title": {
                "text": title,
                "x": 0.02,
                "xanchor": "left",
            },
            "margin": {
                "l": 50,
                "r": 35,
                "t": 55,
                "b": 50,
            },
            "paper_bgcolor": "rgba(0,0,0,0)",
            "plot_bgcolor": "rgba(0,0,0,0)",
            "hovermode": "x unified",
            "legend": {
                "orientation": "h",
                "y": -0.2,
            },
            "font": {
                "size": 12,
            },
        }

        layout.update(kwargs)

        return layout

    def _results(self, processes):
        counts = Counter()

        for process in processes:
            for result in process.batch.test_results.all():
                status = getattr(
                    result,
                    "status",
                    None,
                )

                if status in self.result_statuses:
                    counts[status] += 1

        return {
            status: counts.get(status, 0)
            for status in self.result_statuses
        }

    def _turnaround_values(self, processes):
        values = []

        for process in processes:
            value = self._turnaround_for_process(
                process
            )

            if value is not None and value >= 0:
                values.append(value)

        return values

    def _turnaround_for_process(self, process):
        start = getattr(
            process,
            "analysis_started_at",
            None,
        )

        end = getattr(
            process,
            "completed_at",
            None,
        )

        if not start or not end:
            return None

        return round(
            (end - start).total_seconds() / 3600,
            2,
        )

    def _deadline(self, process):
        request = getattr(
            process,
            "analysis_request",
            None,
        )

        if request is None:
            return None

        return getattr(
            request,
            "deadline",
            None,
        )

    def _is_overdue(self, process):
        deadline = self._deadline(process)

        if deadline is None:
            return False

        if process.status in {
            "COMPLETED",
            "REJECTED",
            "RETURNED",
            "UTILIZED",
        }:
            return False

        return self._now > deadline

    def _chemical_name(self, process):
        chemical = getattr(
            process.batch,
            "chemical",
            None,
        )

        return (
            getattr(chemical, "name", None)
            or getattr(chemical, "code", None)
            or "Unknown"
        )

    def _supplier_name(self, process):
        supplier = getattr(
            process.batch,
            "supplier",
            None,
        )

        return (
            getattr(supplier, "name", None)
            or "Unassigned"
        )

    def _user_name(self, user):
        if not user:
            return "Unassigned"

        full_name = user.get_full_name()

        return (
            full_name.strip()
            or getattr(user, "username", None)
            or "Unknown"
        )

    def _counter_dict(self, value):
        return {
            key: dict(counter)
            for key, counter in value.items()
        }

    def _decimal(self, value):
        if value is None:
            return None

        try:
            return Decimal(str(value))
        except Exception:
            return None

    def _mean(self, values):
        if not values:
            return 0

        return round(
            float(statistics.mean(values)),
            2,
        )

    def _median(self, values):
        if not values:
            return 0

        return round(
            float(statistics.median(values)),
            2,
        )

    def _std(self, values):
        if len(values) < 2:
            return 0

        return round(
            float(statistics.stdev(values)),
            2,
        )

    def _percentile(self, values, percentile):
        if not values:
            return 0

        return round(
            float(
                np.percentile(
                    np.asarray(values),
                    percentile,
                )
            ),
            2,
        )

    def _skew(self, values):
        if len(values) < 3:
            return 0

        if scipy_stats:
            try:
                return round(
                    float(
                        scipy_stats.skew(
                            values,
                            bias=False,
                        )
                    ),
                    2,
                )
            except Exception:
                pass

        return 0

    def _control_limits(self, values):
        if len(values) < 2:
            return {
                "mean": self._mean(values),
                "ucl": None,
                "lcl": None,
            }

        mean = float(np.mean(values))
        std = float(np.std(values))

        return {
            "mean": round(mean, 2),
            "ucl": round(mean + 3 * std, 2),
            "lcl": round(
                max(mean - 3 * std, 0),
                2,
            ),
        }

    def _json_safe(self, value):
        if isinstance(value, dict):
            return {
                str(k): self._json_safe(v)
                for k, v in value.items()
            }

        if isinstance(value, (list, tuple)):
            return [
                self._json_safe(v)
                for v in value
            ]

        if isinstance(value, Decimal):
            return float(value)

        if isinstance(value, np.generic):
            return value.item()

        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None

        return value

    def render_to_response(self, context, **response_kwargs):
        export = self.request.GET.get("export")

        if export:
            return self._export(
                context,
                export,
            )

        return super().render_to_response(
            context,
            **response_kwargs,
        )

    def _export(self, context, export):
        payload = self._json_safe(context)

        if export == "json":
            return JsonResponse(
                payload,
                json_dumps_params={
                    "indent": 2,
                },
            )

        if export == "csv":
            return self._export_csv(
                context
            )

        if export == "xlsx":
            return self._export_xlsx(
                context
            )

        if export == "docx":
            return self._export_docx(
                context
            )

        if export == "pdf":
            return self._export_pdf(
                context
            )

        if export == "zip":
            return self._export_zip(
                context
            )

        return JsonResponse(
            {
                "error": "Unsupported export format."
            },
            status=400,
        )

    def _export_csv(self, context):
        response = HttpResponse(
            content_type="text/csv"
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="chemical_analysis_dashboard.csv"'
        )

        writer = csv.writer(response)

        writer.writerow(
            [
                "Metric",
                "Value",
            ]
        )

        for key, value in context["kpis"].items():
            writer.writerow(
                [
                    key,
                    value,
                ]
            )

        writer.writerow([])
        writer.writerow(
            [
                "Status",
                "Count",
            ]
        )

        for status, data in context[
            "workflow"
        ]["stages"].items():
            writer.writerow(
                [
                    data["label"],
                    data["count"],
                ]
            )

        return response

    def _export_xlsx(self, context):
        output = io.BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl",
        ) as writer:
            kpis = pd.DataFrame(
                [
                    {
                        "metric": key,
                        "value": value,
                    }
                    for key, value in context[
                        "kpis"
                    ].items()
                ]
            )

            kpis.to_excel(
                writer,
                sheet_name="KPIs",
                index=False,
            )

            workflow = pd.DataFrame(
                [
                    {
                        "status": data["code"],
                        "label": data["label"],
                        "count": data["count"],
                    }
                    for data in context[
                        "workflow"
                    ]["stages"].values()
                ]
            )

            workflow.to_excel(
                writer,
                sheet_name="Workflow",
                index=False,
            )

            for dimension, rows in context[
                "comparatives"
            ].items():
                pd.DataFrame(rows).to_excel(
                    writer,
                    sheet_name=dimension[:31],
                    index=False,
                )

            risk = pd.DataFrame(
                context["risk"]
            )

            risk.to_excel(
                writer,
                sheet_name="Risk",
                index=False,
            )

        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="chemical_analysis_dashboard.xlsx"'
        )

        return response

    def _export_docx(self, context):
        from docx import Document

        document = Document()

        document.add_heading(
            "Chemical Analysis Dashboard",
            level=0,
        )

        document.add_paragraph(
            f"Generated: "
            f"{self._now.strftime('%Y-%m-%d %H:%M')}"
        )

        document.add_heading(
            "Key Performance Indicators",
            level=1,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"

        for key, value in context[
            "kpis"
        ].items():
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value)

        document.add_heading(
            "Workflow",
            level=1,
        )

        table = document.add_table(
            rows=1,
            cols=3,
        )

        table.rows[0].cells[0].text = "Status"
        table.rows[0].cells[1].text = "Label"
        table.rows[0].cells[2].text = "Count"

        for data in context[
            "workflow"
        ]["stages"].values():
            cells = table.add_row().cells
            cells[0].text = data["code"]
            cells[1].text = data["label"]
            cells[2].text = str(data["count"])

        document.add_heading(
            "Risk Register",
            level=1,
        )

        for item in context["risk"]:
            document.add_paragraph(
                f"{item['level']}: "
                f"{item['title']} "
                f"({item['count']}) — "
                f"{item['description']}"
            )

        output = io.BytesIO()

        document.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="chemical_analysis_dashboard.docx"'
        )

        return response

    def _export_pdf(self, context):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
            Paragraph,
        )

        output = io.BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
        )

        styles = getSampleStyleSheet()

        elements = [
            Paragraph(
                "Chemical Analysis Dashboard",
                styles["Title"],
            ),
            Spacer(1, 6 * mm),
            Paragraph(
                self._now.strftime(
                    "Generated %Y-%m-%d %H:%M"
                ),
                styles["Normal"],
            ),
            Spacer(1, 5 * mm),
        ]

        rows = [
            ["Metric", "Value"]
        ]

        for key, value in context[
            "kpis"
        ].items():
            rows.append(
                [
                    str(key),
                    str(value),
                ]
            )

        table = Table(rows)

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                ]
            )
        )

        elements.append(table)
        elements.append(
            Spacer(1, 6 * mm)
        )

        elements.append(
            Paragraph(
                "Workflow Status",
                styles["Heading2"],
            )
        )

        rows = [
            [
                "Status",
                "Count",
            ]
        ]

        for data in context[
            "workflow"
        ]["stages"].values():
            rows.append(
                [
                    data["label"],
                    str(data["count"]),
                ]
            )

        table = Table(rows)

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                ]
            )
        )

        elements.append(table)

        document.build(elements)

        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/pdf",
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="chemical_analysis_dashboard.pdf"'
        )

        return response

    def _export_zip(self, context):
        files = {
            "dashboard.json": self._json_export_bytes(
                context
            ),
            "dashboard.csv": self._csv_export_bytes(
                context
            ),
            "dashboard.xlsx": self._xlsx_export_bytes(
                context
            ),
            "dashboard.docx": self._docx_export_bytes(
                context
            ),
            "dashboard.pdf": self._pdf_export_bytes(
                context
            ),
        }

        manifest = {
            "generated_at": self._now.isoformat(),
            "format": "chemical-analysis-dashboard",
            "filters": self._filters,
            "files": list(files),
        }

        files["manifest.json"] = json.dumps(
            manifest,
            indent=2,
            default=str,
        ).encode()

        output = io.BytesIO()

        with zipfile.ZipFile(
            output,
            "w",
            zipfile.ZIP_DEFLATED,
        ) as archive:
            for filename, content in files.items():
                archive.writestr(
                    filename,
                    content,
                )

        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/zip",
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="chemical_analysis_dashboard.zip"'
        )

        return response

    def _json_export_bytes(self, context):
        return json.dumps(
            self._json_safe(context),
            indent=2,
            default=str,
        ).encode()

    def _csv_export_bytes(self, context):
        response = io.StringIO()
        writer = csv.writer(response)

        writer.writerow(
            [
                "Metric",
                "Value",
            ]
        )

        for key, value in context[
            "kpis"
        ].items():
            writer.writerow(
                [
                    key,
                    value,
                ]
            )

        return response.getvalue().encode()

    def _xlsx_export_bytes(self, context):
        output = io.BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl",
        ) as writer:
            pd.DataFrame(
                [
                    {
                        "metric": key,
                        "value": value,
                    }
                    for key, value in context[
                        "kpis"
                    ].items()
                ]
            ).to_excel(
                writer,
                sheet_name="KPIs",
                index=False,
            )

            pd.DataFrame(
                [
                    {
                        "status": data["code"],
                        "label": data["label"],
                        "count": data["count"],
                    }
                    for data in context[
                        "workflow"
                    ]["stages"].values()
                ]
            ).to_excel(
                writer,
                sheet_name="Workflow",
                index=False,
            )

        output.seek(0)

        return output.getvalue()

    def _docx_export_bytes(self, context):
        from docx import Document

        document = Document()

        document.add_heading(
            "Chemical Analysis Dashboard",
            0,
        )

        for key, value in context[
            "kpis"
        ].items():
            document.add_paragraph(
                f"{key}: {value}"
            )

        output = io.BytesIO()
        document.save(output)

        return output.getvalue()

    def _pdf_export_bytes(self, context):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
        )
        from reportlab.lib.styles import getSampleStyleSheet

        output = io.BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=12 * mm,
            leftMargin=12 * mm,
        )

        styles = getSampleStyleSheet()

        elements = [
            Paragraph(
                "Chemical Analysis Dashboard",
                styles["Title"],
            )
        ]

        for key, value in context[
            "kpis"
        ].items():
            elements.append(
                Paragraph(
                    f"{key}: {value}",
                    styles["Normal"],
                )
            )

        document.build(elements)

        output.seek(0)

        return output.getvalue()


class Dashboard3JSONView(Dashboard3View):
    def render_to_response(self, context, **response_kwargs):
        payload = {
            key: value
            for key, value in context.items()
            if key not in {
                "view",
                "object",
                "paginator",
                "page_obj",
                "is_paginated",
            }
        }

        return JsonResponse(
            self._json_safe(payload),
            safe=False,
        )


class Dashboard3RefreshView(Dashboard3View):
    template_name = "ca/dashboard/pages/dashboard3.html"

    def get_template_names(self):
        if (
            self.request.headers.get("HX-Request")
            or self.request.GET.get("partial") == "1"
        ):
            return [
                "ca/dashboard/pages/dashboard3_live.html"
            ]

        return [self.template_name]

