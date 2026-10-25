from __future__ import annotations

import csv
import io
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Count, Exists, OuterRef, Q, Subquery
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from django.views.generic import TemplateView, View
from .mixins import CaWanjalaMixin
from .models import (
    BatchTestResult,
    Chemical,
    ChemicalAnalysisProcess,
    ChemicalAnalysisRequest,
    ChemicalBatch,
    ChemicalLocation,
    ChemicalBatchHold,
    ChemicalInventoryMovement,
    Supplier,
)


BATCH_REGISTER_COLUMNS = {
    "sample": ("Sample ID", "sample_id", "ca.view_chemicalbatch"),
    "batch": ("Supplier batch", "batch_number", "ca.view_chemicalbatch"),
    "chemical": ("Chemical", "chemical", "ca.view_chemicalbatch"),
    "supplier": ("Supplier", "supplier", "ca.view_chemicalbatch"),
    "plant": ("Treatment plant", "plant", "ca.view_chemicalbatch"),
    "location": ("Storage location", "location", "ca.view_chemicalbatch"),
    "received": ("Received", "received_date", "ca.view_chemicalbatch"),
    "age": ("Age", "age", "ca.view_chemicalbatch"),
    "quantity": ("Received quantity", "quantity_received", "ca.view_chemicalbatch"),
    "available": ("Available", "available_quantity", "ca.view_chemicalbatch"),
    "status": ("Batch state", "status", "ca.view_chemicalbatch"),
    "stage": ("Workflow stage", "analysis_stage", "ca.view_chemicalanalysisprocess"),
    "analysts": ("Assigned analysts", "assigned_analysts", "ca.view_chemicalanalysisrequest"),
    "checklist": ("Checklist", "checklist_state", "ca.view_chemicalanalysisrequest"),
    "deadline": ("Analysis deadline", "analysis_deadline", "ca.view_chemicalanalysisrequest"),
    "sla": ("SLA signal", "sla_state", "ca.view_chemicalanalysisrequest"),
    "test_count": ("Tests recorded", "test_count", "ca.view_batchtestresult"),
    "result": ("Latest result", "latest_result", "ca.view_batchtestresult"),
    "result_date": ("Last tested", "latest_tested_at", "ca.view_batchtestresult"),
    "coa": ("CoA", "coa_state", "ca.view_chemicalcoa"),
    "hold": ("Intake hold", "hold_state", "ca.view_chemicalbatch"),
    "received_by": ("Received by", "received_by", "ca.view_chemicalbatch"),
    "updated": ("Last updated", "updated_at", "ca.view_chemicalbatch"),
}

BATCH_REGISTER_SORTS = {
    "received_date": "Received date",
    "-received_date": "Newest receipt",
    "updated_at": "Last updated",
    "-updated_at": "Recently updated",
    "sample_id": "Sample ID A–Z",
    "-sample_id": "Sample ID Z–A",
    "batch_number": "Batch number A–Z",
    "-batch_number": "Batch number Z–A",
    "chemical__name": "Chemical A–Z",
    "-chemical__name": "Chemical Z–A",
    "supplier__name": "Supplier A–Z",
    "-supplier__name": "Supplier Z–A",
    "analysis_request__analysis_deadline": "Earliest deadline",
    "-analysis_request__analysis_deadline": "Latest deadline",
    "quantity_received": "Quantity low to high",
    "-quantity_received": "Quantity high to low",
    "status": "State A–Z",
    "-status": "State Z–A",
}


def _valid_date(value):
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


class ChemicalBatchRegisterQueryMixin:
    """Shared, plant-scoped query and filter contract for the batch register."""

    def get_batch_register_queryset(self, *, include_workflow_tasks=True):
        queryset = self.scope_queryset(
            ChemicalBatch.objects.all(), field="treatment_plant_id"
        )
        latest_inventory = ChemicalInventoryMovement.objects.filter(
            batch_id=OuterRef("pk")
        ).order_by("-created_at", "-pk")
        latest_result = BatchTestResult.objects.filter(
            batch_id=OuterRef("pk")
        ).order_by("-tested_at", "-pk")
        result_count = (
            BatchTestResult.objects.filter(batch_id=OuterRef("pk"))
            .order_by().values("batch_id").annotate(total=Count("pk"))
            .values("total")[:1]
        )
        open_holds = ChemicalBatchHold.objects.filter(
            batch_id=OuterRef("pk"), status=ChemicalBatchHold.Status.OPEN
        )
        queryset = queryset.annotate(
            register_inventory_balance=Subquery(
                latest_inventory.values("quantity_after")[:1]
            ),
            register_latest_result=Subquery(latest_result.values("status")[:1]),
            register_latest_tested_at=Subquery(latest_result.values("tested_at")[:1]),
            register_test_count=Subquery(result_count),
            register_has_open_hold=Exists(open_holds),
        ).select_related(
            "chemical", "supplier", "location", "treatment_plant", "received_by",
            "analysis_request", "analysis_process", "coa",
        )
        prefetches = ["analysis_request__assigned_analysts", "analysis_request__checklists"]
        if include_workflow_tasks:
            prefetches.append("analysis_process__task_set")
        queryset = queryset.prefetch_related(*prefetches)

        params = self.request.GET
        query = (params.get("q") or "").strip()
        if query:
            queryset = queryset.filter(
                Q(sample_id__icontains=query)
                | Q(batch_number__icontains=query)
                | Q(delivery_note_number__icontains=query)
                | Q(chemical__name__icontains=query)
                | Q(chemical__code__icontains=query)
                | Q(supplier__name__icontains=query)
                | Q(location__name__icontains=query)
                | Q(treatment_plant__name__icontains=query)
            )

        for parameter, field in (
            ("status", "status"),
            ("chemical", "chemical_id"),
            ("supplier", "supplier_id"),
            ("location", "location_id"),
        ):
            value = (params.get(parameter) or "").strip()
            if value and value.isdigit():
                queryset = queryset.filter(**{field: int(value)})
            elif parameter == "status" and value in dict(ChemicalBatch.Status.choices):
                queryset = queryset.filter(status=value)

        received_from = _valid_date(params.get("received_from"))
        received_to = _valid_date(params.get("received_to"))
        if received_from:
            queryset = queryset.filter(received_date__gte=received_from)
        if received_to:
            queryset = queryset.filter(received_date__lte=received_to)

        for parameter, lookup in (("qty_min", "gte"), ("qty_max", "lte")):
            try:
                quantity = Decimal(str(params.get(parameter, "")).strip())
                if quantity.is_finite() and quantity >= 0:
                    queryset = queryset.filter(**{f"quantity_received__{lookup}": quantity})
            except (InvalidOperation, TypeError, ValueError):
                pass

        can_view_coa = self.is_ca_superuser() or self.request.user.has_perm("ca.view_chemicalcoa")
        coa_filter = params.get("coa")
        if can_view_coa and coa_filter in {"ready", "pending"}:
            queryset = queryset.filter(coa_generated=(coa_filter == "ready"))

        hold_filter = params.get("hold")
        if hold_filter in {"open", "clear"}:
            queryset = queryset.filter(register_has_open_hold=(hold_filter == "open"))

        can_view_processes = self.is_ca_superuser() or self.request.user.has_perm("ca.view_chemicalanalysisprocess")
        stage = (params.get("stage") or "").strip()
        if can_view_processes:
            if stage in dict(ChemicalAnalysisProcess.WorkflowStatus.choices):
                queryset = queryset.filter(analysis_process__status=stage)
            elif stage == "NO_PROCESS":
                queryset = queryset.filter(analysis_process__isnull=True)

        can_view_requests = self.is_ca_superuser() or self.request.user.has_perm("ca.view_chemicalanalysisrequest")
        analyst = (params.get("analyst") or "").strip()
        if can_view_requests and analyst.isdigit():
            queryset = queryset.filter(analysis_request__assigned_analysts__pk=int(analyst))
        sla = params.get("sla")
        if can_view_requests and sla in {"overdue", "due_soon", "on_track", "complete", "no_deadline", "cancelled", "no_request"}:
            now = timezone.now()
            active_request = Q(
                analysis_request__status__in=(
                    ChemicalAnalysisRequest.Status.PENDING,
                    ChemicalAnalysisRequest.Status.IN_PROGRESS,
                ),
                analysis_request__completed_at__isnull=True,
            )
            if sla == "overdue":
                queryset = queryset.filter(active_request, analysis_request__analysis_deadline__lt=now)
            elif sla == "due_soon":
                queryset = queryset.filter(
                    active_request,
                    analysis_request__analysis_deadline__gte=now,
                    analysis_request__analysis_deadline__lte=now + timedelta(hours=24),
                )
            elif sla == "on_track":
                queryset = queryset.filter(
                    active_request, analysis_request__analysis_deadline__gt=now + timedelta(hours=24)
                )
            elif sla == "complete":
                queryset = queryset.filter(analysis_request__completed_at__isnull=False)
            elif sla == "cancelled":
                queryset = queryset.filter(analysis_request__status=ChemicalAnalysisRequest.Status.CANCELLED)
            elif sla == "no_request":
                queryset = queryset.filter(analysis_request__isnull=True)
            else:
                queryset = queryset.filter(
                    active_request,
                    analysis_request__isnull=False,
                    analysis_request__analysis_deadline__isnull=True,
                )

        age = params.get("age")
        today = timezone.localdate()
        if age == "recent":
            queryset = queryset.filter(received_date__gte=today - timedelta(days=30))
        elif age == "30_90":
            queryset = queryset.filter(
                received_date__lt=today - timedelta(days=30),
                received_date__gte=today - timedelta(days=90),
            )
        elif age == "90_plus":
            queryset = queryset.filter(received_date__lt=today - timedelta(days=90))

        sort = params.get("sort", "-received_date")
        if sort not in BATCH_REGISTER_SORTS:
            sort = "-received_date"
        return queryset.order_by(sort, "-pk").distinct()

    def get_allowed_register_columns(self):
        return {
            key: definition
            for key, definition in BATCH_REGISTER_COLUMNS.items()
            if self.is_ca_superuser() or self.request.user.has_perm(definition[2])
        }


def _enrich_batch_register_rows(
    batches,
    user,
    can_view_requests,
    can_view_processes,
    is_ca_superuser=False,
    include_actions=True,
):
    today = timezone.localdate()
    for batch in batches:
        batch.days_in_store = max(0, (today - batch.received_date).days) if batch.received_date else None
        if batch.register_inventory_balance is not None:
            batch.register_available_quantity = batch.register_inventory_balance
        else:
            batch.register_available_quantity = max(
                Decimal("0"),
                Decimal(str(batch.quantity_received or 0))
                - Decimal(str(batch.quantity_utilized or 0))
                - Decimal(str(batch.quantity_returned or 0)),
            )
        batch.latest_result = batch.register_latest_result or "—"
        batch.latest_tested_at = batch.register_latest_tested_at
        batch.test_count = batch.register_test_count or 0
        batch.has_open_hold = batch.register_has_open_hold
        batch.hold_state = "Open" if batch.has_open_hold else "Clear"
        batch.coa_state = "Generated" if batch.coa_generated else "Pending"
        batch.checklist_state = "—"
        batch.assigned_analysts_label = "—"
        batch.analysis_deadline = None
        batch.sla_state = "—"
        batch.sla_state_key = "none"
        batch.analysis_stage_label = "Not started"
        batch.analysis_stage_key = "not-started"
        if can_view_processes:
            process = getattr(batch, "analysis_process", None)
            if process:
                batch.analysis_stage_label = dict(
                    ChemicalAnalysisProcess.WorkflowStatus.choices
                ).get(str(process.status), str(process.status).replace("_", " ").title())
                batch.analysis_stage_key = str(process.status).lower()
        if can_view_requests:
            request = getattr(batch, "analysis_request", None)
            if request:
                analysts = list(request.assigned_analysts.all())
                batch.assigned_analysts_label = ", ".join(
                    getattr(person, "get_full_name", lambda: "")() or person.get_username()
                    for person in analysts
                ) or "Unassigned"
                checklists = list(request.checklists.all())
                if checklists:
                    completed = sum(bool(item.completed) for item in checklists)
                    batch.checklist_state = f"{completed}/{len(checklists)} complete"
                else:
                    batch.checklist_state = "Not started"
                batch.analysis_deadline = request.analysis_deadline
                now = timezone.now()
                if request.status == ChemicalAnalysisRequest.Status.CANCELLED:
                    batch.sla_state = "Cancelled"
                    batch.sla_state_key = "cancelled"
                elif request.completed_at:
                    batch.sla_state = (
                        "Late" if request.analysis_deadline and request.completed_at > request.analysis_deadline
                        else "Met"
                    )
                    batch.sla_state_key = batch.sla_state.lower()
                elif not request.analysis_deadline:
                    batch.sla_state = "No deadline"
                    batch.sla_state_key = "no-deadline"
                elif request.analysis_deadline < now:
                    batch.sla_state = "Overdue"
                    batch.sla_state_key = "overdue"
                elif request.analysis_deadline <= now + timedelta(hours=24):
                    batch.sla_state = "Due <24h"
                    batch.sla_state_key = "due-soon"
                else:
                    batch.sla_state = "On track"
                    batch.sla_state_key = "on-track"
            else:
                batch.analysis_deadline = None
                batch.sla_state = "No request"
                batch.sla_state_key = "no-request"
        batch.workflow_action_url = ""
        batch.workflow_action_label = ""
        if include_actions and can_view_requests and can_view_processes:
            _set_batch_workflow_action(batch, user, is_ca_superuser=is_ca_superuser)
        if user.has_perm("ca.change_chemicalbatch") or is_ca_superuser or getattr(user, "is_superuser", False):
            batch.edit_url = reverse("ca:batch-update", kwargs={"pk": batch.pk})
        else:
            batch.edit_url = ""
        batch.detail_url = reverse("ca:batch-detail", kwargs={"pk": batch.pk})
    return batches


def _set_batch_workflow_action(batch, user, is_ca_superuser=False):
    request = getattr(batch, "analysis_request", None)
    process = getattr(batch, "analysis_process", None)
    if not request or not process or batch.has_open_hold:
        return
    try:
        if not request.is_assigned_analyst(user) or not process.is_authorized_analyst(user):
            return
    except Exception:
        return
    checklists = list(request.checklists.all())
    if not any(item.completed for item in checklists):
        node_name, label, permission = "checklist", "Complete checklist", "ca.complete_receiving_checklist"
    elif process.status in {
        ChemicalAnalysisProcess.WorkflowStatus.ASSIGNED,
        ChemicalAnalysisProcess.WorkflowStatus.CHECKLIST,
    }:
        node_name, label, permission = "start_analysis", "Start analysis", "ca.start_chemical_analysis"
    elif process.status == ChemicalAnalysisProcess.WorkflowStatus.ANALYSIS:
        node_name, label, permission = "record_result", "Continue analysis", "ca.record_chemical_result"
    else:
        return
    if not (is_ca_superuser or getattr(user, "is_superuser", False)) and not user.has_perm(permission):
        return
    active = {"NEW", "ASSIGNED", "STARTED"}
    for task in process.task_set.all():
        node = task.flow_task
        if not node or node.name != node_name or task.status not in active:
            continue
        try:
            can_open = (
                node.can_assign(user, task) if task.status == "NEW"
                else True if task.status == "ASSIGNED"
                else task.owner_id == user.pk
            )
        except Exception:
            can_open = False
        if can_open:
            batch.workflow_action_url = reverse(
                "ca:batch-workflow-action", kwargs={"pk": batch.pk, "node_name": node_name}
            )
            batch.workflow_action_label = label
            return


class ChemicalAnalysisDashboardView(CaWanjalaMixin, ChemicalBatchRegisterQueryMixin, TemplateView):
    """Compact, permission-aware chemical batch command register."""

    template_name = "ca/command_center.html"

    def _allowed(self, permission: str) -> bool:
        return self.is_ca_superuser() or self.request.user.has_perm(permission)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        can_view_batches = self._allowed("ca.view_chemicalbatch")
        can_view_requests = self._allowed("ca.view_chemicalanalysisrequest")
        can_view_processes = self._allowed("ca.view_chemicalanalysisprocess")
        can_view_results = self._allowed("ca.view_batchtestresult")
        can_view_coa = self._allowed("ca.view_chemicalcoa")
        can_receive = self._allowed("ca.receive_chemical_batch")
        plant_context = context.get("plant_context", {})
        chemical_locations = ChemicalLocation.objects.select_related(
            "treatment_plant"
        ).order_by("treatment_plant__name", "name")
        if self.is_ca_superuser():
            selected_plant_id = plant_context.get("plant_id")
            if selected_plant_id is not None:
                chemical_locations = chemical_locations.filter(
                    treatment_plant_id=selected_plant_id
                )
        else:
            assigned_plant_id = getattr(user, "assigned_treatment_plant_id", None)
            chemical_locations = (
                chemical_locations.filter(treatment_plant_id=assigned_plant_id)
                if assigned_plant_id is not None
                else chemical_locations.none()
            )
        allowed_columns = self.get_allowed_register_columns() if can_view_batches else {}
        selected_columns = [
            key for key in (self.request.GET.get("columns", "").split(","))
            if key in allowed_columns
        ]
        if not selected_columns:
            default_columns = [
                "sample", "chemical", "supplier", "location", "received", "age",
                "quantity", "available", "status", "stage", "sla", "hold",
            ]
            if self.is_ca_superuser():
                default_columns.append("plant")
            selected_columns = [
                key for key in default_columns if key in allowed_columns
            ]

        batches = []
        page_obj = None
        page_numbers = []
        total_batches = 0
        filter_options = {"chemicals": [], "suppliers": [], "locations": [], "analysts": []}
        if can_view_batches:
            base = self.scope_queryset(ChemicalBatch.objects.all(), field="treatment_plant_id")
            filter_options = {
                "chemicals": Chemical.objects.filter(
                    pk__in=base.order_by().values_list("chemical_id", flat=True).distinct()
                ).order_by("name"),
                "suppliers": Supplier.objects.filter(
                    pk__in=base.order_by().values_list("supplier_id", flat=True).distinct()
                ).order_by("name"),
                "locations": ChemicalLocation.objects.filter(
                    pk__in=base.order_by().values_list("location_id", flat=True).distinct()
                ).order_by("name"),
            }
            if can_view_requests:
                from django.contrib.auth import get_user_model
                User = get_user_model()
                filter_options["analysts"] = User.objects.filter(
                    assigned_requests__batch__in=base
                ).distinct().order_by("first_name", "last_name", "username")

            queryset = self.get_batch_register_queryset()
            total_batches = queryset.count()
            try:
                page_size = int(self.request.GET.get("page_size", "50"))
            except (TypeError, ValueError):
                page_size = 50
            if page_size not in {25, 50, 100, 200}:
                page_size = 50
            paginator = Paginator(queryset, page_size)
            try:
                page_obj = paginator.page(self.request.GET.get("page", 1))
            except (PageNotAnInteger, EmptyPage):
                page_obj = paginator.page(1 if not str(self.request.GET.get("page", "1")).isdigit() else paginator.num_pages or 1)
            page_numbers = paginator.get_elided_page_range(
                page_obj.number, on_each_side=1, on_ends=1
            )
            batches = _enrich_batch_register_rows(
                list(page_obj.object_list), user, can_view_requests, can_view_processes,
                is_ca_superuser=self.is_ca_superuser(),
            )

        pagination_params = self.request.GET.copy()
        pagination_params.pop("page", None)
        context.update(
            {
                "dashboard_now": timezone.now(),
                "can_view_batches": can_view_batches,
                "can_view_requests": can_view_requests,
                "can_view_processes": can_view_processes,
                "can_view_results": can_view_results,
                "can_view_coa": can_view_coa,
                "can_receive_batch": can_receive,
                "is_ca_superuser": self.is_ca_superuser(),
                "chemicals": list(Chemical.objects.order_by("name")) if can_receive else [],
                "suppliers": list(Supplier.objects.order_by("name")) if can_receive else [],
                "chemical_locations": list(chemical_locations) if can_receive else [],
                "can_export_batches": can_view_batches,
                "register_batches": batches,
                "register_page": page_obj,
                "register_page_numbers": page_numbers,
                "register_total": total_batches,
                "register_columns": allowed_columns,
                "register_selected_columns": selected_columns,
                "register_filters": self.request.GET,
                "register_filter_options": filter_options,
                "register_sorts": BATCH_REGISTER_SORTS,
                "register_query_without_page": pagination_params.urlencode(),
                "register_page_size": int(self.request.GET.get("page_size", "50"))
                if str(self.request.GET.get("page_size", "50")).isdigit()
                and int(self.request.GET.get("page_size", "50")) in {25, 50, 100, 200}
                else 50,
                "chemical_batch_statuses": ChemicalBatch.Status.choices,
                "workflow_statuses": ChemicalAnalysisProcess.WorkflowStatus.choices,
                "age_buckets": (
                    ("recent", "0–30 days"),
                    ("30_90", "31–90 days"),
                    ("90_plus", "Over 90 days"),
                ),
                "sla_filters": (
                    ("overdue", "Overdue"),
                    ("due_soon", "Due within 24 hours"),
                    ("on_track", "On track"),
                    ("complete", "Completed"),
                    ("no_deadline", "No deadline"),
                    ("cancelled", "Cancelled"),
                    ("no_request", "No analysis request"),
                ),
            }
        )

        if can_receive:
            context["receive_batch_url"] = reverse("ca:batch-create")
        return context


class ChemicalBatchRegisterExportView(
    CaWanjalaMixin, ChemicalBatchRegisterQueryMixin, View
):
    """Exports the current plant-scoped, filtered batch register."""

    permission_required_view = "ca.view_chemicalbatch"

    def get(self, request, *args, **kwargs):
        allowed = self.get_allowed_register_columns()
        requested = [key for key in request.GET.get("columns", "").split(",") if key in allowed]
        if not requested:
            requested = list(allowed)
        labels = [allowed[key][0] for key in requested]
        batches = list(_enrich_batch_register_rows(
            self.get_batch_register_queryset(include_workflow_tasks=False), request.user,
            request.user.has_perm("ca.view_chemicalanalysisrequest") or self.is_ca_superuser(),
            request.user.has_perm("ca.view_chemicalanalysisprocess") or self.is_ca_superuser(),
            is_ca_superuser=self.is_ca_superuser(),
            include_actions=False,
        ))
        records = [_batch_register_export_row(batch, requested) for batch in batches]
        fmt = (request.GET.get("format") or "csv").lower()
        if fmt == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(labels)
            writer.writerows(records)
            response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
            response["Content-Disposition"] = 'attachment; filename="chemical-batch-register.csv"'
            return response
        if fmt in {"xlsx", "excel"}:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Chemical batches"
            sheet.append(labels)
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="10263F")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
            for row in records:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(34, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
            output = io.BytesIO()
            workbook.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = 'attachment; filename="chemical-batch-register.xlsx"'
            return response
        if fmt in {"doc", "docx", "word"}:
            from docx import Document
            from docx.enum.section import WD_ORIENT
            document = Document()
            section = document.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
            document.add_heading("Chemical batch register", 0)
            document.add_paragraph(f"Generated {timezone.localtime():%Y-%m-%d %H:%M} · {len(records)} batches")
            table = document.add_table(rows=1, cols=len(labels))
            table.style = "Light Shading Accent 1"
            for cell, label in zip(table.rows[0].cells, labels):
                cell.text = label
            for row in records:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value or "—")
            output = io.BytesIO()
            document.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            response["Content-Disposition"] = 'attachment; filename="chemical-batch-register.docx"'
            return response
        if fmt == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
            from xml.sax.saxutils import escape
            output = io.BytesIO()
            document = SimpleDocTemplate(output, pagesize=landscape(letter), leftMargin=0.35 * inch, rightMargin=0.35 * inch)
            styles = getSampleStyleSheet()
            data = [[Paragraph(f"<b>{escape(label)}</b>", styles["BodyText"]) for label in labels]]
            data.extend([[Paragraph(escape(str(value or "—")), styles["BodyText"]) for value in row] for row in records])
            printable_width = landscape(letter)[0] - 0.7 * inch
            table = Table(
                data,
                colWidths=[printable_width / max(1, len(labels))] * len(labels),
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10263F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d8e2ed")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f7fb")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            document.build([Paragraph("Chemical batch register", styles["Title"]), Spacer(1, 8), table])
            response = HttpResponse(output.getvalue(), content_type="application/pdf")
            response["Content-Disposition"] = 'attachment; filename="chemical-batch-register.pdf"'
            return response
        return HttpResponse("Unsupported export format.", status=400, content_type="text/plain")


def _batch_register_export_row(batch, columns):
    values = {
        "sample": batch.sample_id,
        "batch": batch.batch_number,
        "chemical": batch.chemical.name if batch.chemical_id else "—",
        "supplier": batch.supplier.name if batch.supplier_id else "—",
        "plant": batch.treatment_plant.name if batch.treatment_plant_id else "—",
        "location": batch.location.name if batch.location_id else "—",
        "received": batch.received_date,
        "age": batch.days_in_store,
        "quantity": f"{batch.quantity_received:g} {batch.unit}",
        "available": f"{batch.register_available_quantity:g} {batch.unit}",
        "status": batch.get_status_display(),
        "stage": batch.analysis_stage_label,
        "analysts": batch.assigned_analysts_label,
        "checklist": batch.checklist_state,
        "deadline": batch.analysis_deadline,
        "sla": batch.sla_state,
        "test_count": batch.test_count,
        "result": batch.latest_result,
        "result_date": batch.latest_tested_at,
        "coa": batch.coa_state,
        "hold": batch.hold_state,
        "received_by": str(batch.received_by) if batch.received_by_id else "—",
        "updated": batch.updated_at,
    }
    return [values[key] for key in columns]
